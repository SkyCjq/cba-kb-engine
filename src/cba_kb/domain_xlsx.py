"""Six-tab candidate export for subsequent native Sheet review and integration."""
from pathlib import Path
import json
from openpyxl import Workbook

from .registration_domain import TABLES
from .xlsx import normalize_archive, STAMP, CREATOR

SHEET_NAMES = {
    'domestic_registrations': 'domestic_registrations',
    'domestic_transaction_windows': 'domestic_transaction_windows',
    'registration_status_events': 'registration_status_events',
    'foreign_registration_snapshots': 'foreign_registration_snapshots',
    'foreign_priority_right_snapshots': 'foreign_right_snapshots',
    'foreign_priority_right_transactions': 'foreign_right_transactions',
}


def write(path, records):
    path = Path(path)
    if path.exists():
        raise ValueError('Candidate workbook path must be new')
    book = Workbook()
    book.properties.created = STAMP
    book.properties.modified = STAMP
    book.properties.creator = CREATOR
    book.properties.lastModifiedBy = CREATOR
    first = True
    for name, headers in TABLES.items():
        page = book.active if first else book.create_sheet()
        first = False
        page.title = SHEET_NAMES[name]
        page.append(list(headers))
        for index, row in enumerate(records.get(name) or [], 1):
            if sorted(row) != sorted(headers):
                raise ValueError(f'{name} row {index} does not match schema')
            values = []
            for header in headers:
                value = row.get(header)
                if isinstance(value, (list, tuple, dict)):
                    value = json.dumps(value, ensure_ascii=False, sort_keys=True)
                values.append(value)
            page.append(values)
            for cell in page[page.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return normalize_archive(path)
