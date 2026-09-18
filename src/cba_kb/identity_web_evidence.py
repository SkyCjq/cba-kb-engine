"""Official web evidence enrichment and scalable human review support.

This module never creates final identity authority. It produces candidate-only
evidence metadata, explicit human-authority expansion, and non-canonical XLSX
working projections.
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import time
import unicodedata
import urllib.error
import urllib.request
from collections import defaultdict
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection

from .evidence_ledger import canonical_bytes
from .identity_coverage import (
    HUMAN_REVIEW_FIELDS,
    MACHINE_REVIEW_FIELDS,
    REVIEW_FIELDS,
    validate_review_packet,
)
from .player_identity import validate_player_uid


SCHEMA_VERSION = 1
EXTRACTOR_VERSION = "identity-web-evidence-v1"
EVIDENCE_CLASSES = frozenset({
    "W1_OFFICIAL_PERSON_ID_EXACT",
    "W2_OFFICIAL_BIO_MULTI_SOURCE",
    "W3_OFFICIAL_CONTINUITY",
    "W4_DISCOVERY_SUPPORT",
    "WX_CONFLICT",
})
BATCH_PREDICATES = frozenset({
    "BATCH_EXISTING_W1",
    "BATCH_EXISTING_W2",
    "BATCH_NO_SAFE_SEARCH_EXHAUSTED",
})
WORKBOOK_SHEETS = (
    "Review",
    "Groups",
    "Batches",
    "Evidence_Index",
    "Instructions",
)
INSTRUCTIONS_ROWS = (
    ("Authority boundary",),
    ("Workbook is non-canonical human workflow projection only.",),
    ("Only human-owned cells may be edited.",),
    ("Do not paste source bodies into this workbook.",),
)
SOURCE_TIERS = {
    "A0": frozenset({
        "cba.net.cn",
        "www.cba.net.cn",
        "cbanetcdn.cba.net.cn",
    }),
    "A1": frozenset({
        "cbaleague.com",
        "www.cbaleague.com",
        "image.cbaleague.com",
    }),
    "A2": frozenset({
        "cbaleague.com",
        "www.cbaleague.com",
    }),
}
DISCOVERY_PROVIDERS = frozenset({
    "search_engine",
    "mainstream_media",
    "encyclopedia",
    "forum",
    "social_post",
    "search_snippet",
    "agent",
})
CLAIM_TYPES = frozenset({
    "OFFICIAL_PLAYER_NAME",
    "OFFICIAL_BIRTH_DATE",
    "OFFICIAL_REGISTRATION_UNIT",
    "OFFICIAL_SEASON",
    "OFFICIAL_TEAM",
    "OFFICIAL_JERSEY_NUMBER",
    "OFFICIAL_SOURCE_DECLARED_PERSON_ID",
    "OFFICIAL_TRANSACTION_OR_REGISTRATION_EVENT",
})
EXTERNAL_ID_NAMESPACE = re.compile(r"[A-Z][A-Z0-9_]*\Z")
_SHA256 = re.compile(r"[0-9a-f]{64}\Z")
_DATE = re.compile(r"\d{4}-\d{2}-\d{2}\Z")
_SUSPICIOUS_TEXT = re.compile(r"[�]|Ã|Â|锟|拷|娌|鎹")
_DOMAIN_LAST_REQUEST = {}


class IdentityWebEvidenceError(RuntimeError):
    pass


def _required_text(value, label):
    if not isinstance(value, str) or not value.strip():
        raise IdentityWebEvidenceError(f"{label}_REQUIRED")
    return value.strip()


def _required_sha256(value, label):
    value = _required_text(value, label)
    if not _SHA256.fullmatch(value):
        raise IdentityWebEvidenceError(f"{label}_INVALID")
    return value


def _normalize_text(value):
    text = unicodedata.normalize("NFKC", _required_text(value, "VALUE"))
    return " ".join(text.strip().split())


def source_domain(url):
    parsed = urlparse(_required_text(url, "SOURCE_URL"))
    if parsed.scheme not in {"http", "https"}:
        raise IdentityWebEvidenceError("PUBLIC_HTTP_URL_REQUIRED")
    if parsed.username or parsed.password:
        raise IdentityWebEvidenceError("URL_CREDENTIALS_FORBIDDEN")
    if not parsed.hostname:
        raise IdentityWebEvidenceError("SOURCE_DOMAIN_REQUIRED")
    return parsed.hostname.lower()


def validate_source_url(
    url,
    *,
    source_tier,
    private_b0_domains=None,
):
    domain = source_domain(url)
    parsed = urlparse(url)
    path = parsed.path.lower()
    if any(token in path for token in (
        "/login",
        "/signin",
        "/oauth",
        "/private-api",
        "/api/private",
    )):
        raise IdentityWebEvidenceError("AUTHENTICATED_RESOURCE_FORBIDDEN")
    if source_tier in SOURCE_TIERS:
        if domain not in SOURCE_TIERS[source_tier]:
            raise IdentityWebEvidenceError("SOURCE_DOMAIN_NOT_ALLOWED")
    elif source_tier == "B0":
        approved = {
            value.lower()
            for value in (private_b0_domains or [])
        }
        if domain not in approved:
            raise IdentityWebEvidenceError("B0_DOMAIN_NOT_APPROVED")
    elif source_tier == "D0":
        raise IdentityWebEvidenceError("DISCOVERY_SOURCE_NOT_AUTHORITY")
    else:
        raise IdentityWebEvidenceError("SOURCE_TIER_INVALID")
    return {
        "url": url,
        "domain": domain,
        "source_tier": source_tier,
    }


def validate_redirect(
    original_url,
    destination_url,
    *,
    source_tier,
    private_b0_domains=None,
):
    validate_source_url(
        original_url,
        source_tier=source_tier,
        private_b0_domains=private_b0_domains,
    )
    return validate_source_url(
        destination_url,
        source_tier=source_tier,
        private_b0_domains=private_b0_domains,
    )


def validate_fetch_request(url, *, method="GET", headers=None):
    if method not in {"GET", "HEAD"}:
        raise IdentityWebEvidenceError("PUBLIC_FETCH_METHOD_INVALID")
    headers = headers or {}
    if any(str(key).lower() in {"authorization", "cookie"} for key in headers):
        raise IdentityWebEvidenceError("AUTHENTICATED_HEADERS_FORBIDDEN")
    parsed = urlparse(_required_text(url, "SOURCE_URL"))
    query = parsed.query.lower()
    blocked_query_keys = ("token", "session", "auth", "cookie")
    if any(f"{token}=" in query for token in blocked_query_keys):
        raise IdentityWebEvidenceError("AUTHENTICATED_QUERY_FORBIDDEN")
    source_domain(url)
    return {"url": url, "method": method, "headers": {}}


def build_discovery_record(
    *,
    query,
    discovered_url,
    discovery_provider,
    discovered_at,
    title_hint=None,
):
    if discovery_provider not in DISCOVERY_PROVIDERS:
        raise IdentityWebEvidenceError("DISCOVERY_PROVIDER_INVALID")
    return {
        "schema_version": SCHEMA_VERSION,
        "authority": "DISCOVERY_ONLY",
        "query": _required_text(query, "QUERY"),
        "discovered_url": discovered_url,
        "discovered_domain": source_domain(discovered_url),
        "discovered_at": _required_text(discovered_at, "DISCOVERED_AT"),
        "discovery_provider": discovery_provider,
        "title_hint": title_hint,
    }


def validate_discovery_record(value):
    if not isinstance(value, dict):
        raise IdentityWebEvidenceError("DISCOVERY_OBJECT_REQUIRED")
    required = {
        "schema_version",
        "authority",
        "query",
        "discovered_url",
        "discovered_domain",
        "discovered_at",
        "discovery_provider",
        "title_hint",
    }
    if set(value) != required:
        raise IdentityWebEvidenceError("DISCOVERY_SCHEMA_INVALID")
    if value["authority"] != "DISCOVERY_ONLY":
        raise IdentityWebEvidenceError("DISCOVERY_AUTHORITY_INVALID")
    if value["discovery_provider"] not in DISCOVERY_PROVIDERS:
        raise IdentityWebEvidenceError("DISCOVERY_PROVIDER_INVALID")
    if source_domain(value["discovered_url"]) != value["discovered_domain"]:
        raise IdentityWebEvidenceError("DISCOVERY_DOMAIN_MISMATCH")
    return value


def collect_evidence_from_responses(
    responses,
    *,
    output_root,
    extracted_at,
    source_tier,
    source_kind,
    extractor_version=EXTRACTOR_VERSION,
    max_bytes=10_000_000,
    timeout_seconds=10,
    private_b0_domains=None,
):
    output_root = Path(output_root)
    raw_root = output_root / "raw"
    raw_root.mkdir(parents=True, exist_ok=True)
    items = []
    for response in responses:
        discovery = validate_discovery_record(response["discovery"])
        body = response["body"]
        if not isinstance(body, bytes):
            raise IdentityWebEvidenceError("RESPONSE_BODY_BYTES_REQUIRED")
        validate_response_guard(
            status=response["status"],
            content_type=response["content_type"],
            content=body,
            max_bytes=max_bytes,
            timeout_seconds=timeout_seconds,
        )
        validate_source_url(
            discovery["discovered_url"],
            source_tier=source_tier,
            private_b0_domains=private_b0_domains,
        )
        claims = response.get("claims")
        if claims is None:
            if response["content_type"].startswith("text/html"):
                claims = extract_html_claims(body)
            elif response["content_type"] == "application/pdf":
                claims = extract_pdf_claims(
                    body,
                    text_extractor=response.get("text_extractor"),
                )
            else:
                claims = []
        snapshot_name = content_sha256(body) + ".bin"
        snapshot_path = raw_root / snapshot_name
        if snapshot_path.exists():
            if snapshot_path.read_bytes() != body:
                raise IdentityWebEvidenceError("RAW_SNAPSHOT_CONFLICT")
        else:
            snapshot_path.write_bytes(body)
        items.append(build_evidence_item(
            source_tier=source_tier,
            source_kind=source_kind,
            source_url=discovery["discovered_url"],
            fetched_at=extracted_at,
            http_status=response["status"],
            content_type=response["content_type"],
            content=body,
            raw_snapshot_path=str(snapshot_path),
            claims=claims,
            record_keys=response.get("record_keys", []),
            published_at=response.get("published_at"),
            extraction_warnings=response.get("extraction_warnings"),
            extractor_version=extractor_version,
            private_b0_domains=private_b0_domains,
        ))
    return build_evidence_manifest(items)


def validate_response_guard(
    *,
    status,
    content_type,
    content,
    max_bytes,
    timeout_seconds,
):
    if timeout_seconds <= 0:
        raise IdentityWebEvidenceError("FETCH_TIMEOUT_REQUIRED")
    if status < 200 or status >= 300:
        raise IdentityWebEvidenceError("FETCH_STATUS_FAILED")
    if len(content) > max_bytes:
        raise IdentityWebEvidenceError("FETCH_RESPONSE_TOO_LARGE")
    if not content_type:
        raise IdentityWebEvidenceError("CONTENT_TYPE_REQUIRED")
    return {
        "status": status,
        "content_type": content_type,
        "bytes": len(content),
    }


def _urllib_transport(url, *, method, timeout, headers):
    request = urllib.request.Request(url, method=method, headers=headers)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return {
                "status": response.status,
                "content_type": response.headers.get_content_type(),
                "body": response.read(),
                "redirects": [],
            }
    except urllib.error.HTTPError as exc:
        return {
            "status": exc.code,
            "content_type": exc.headers.get_content_type(),
            "body": exc.read(),
            "redirects": [],
        }


def fetch_official_resource(
    url,
    *,
    source_tier,
    method="GET",
    headers=None,
    timeout_seconds=10,
    max_bytes=10_000_000,
    max_retries=2,
    per_domain_interval=1.0,
    transport=None,
    sleep=time.sleep,
    clock=time.monotonic,
    private_b0_domains=None,
):
    request = validate_fetch_request(url, method=method, headers=headers)
    validate_source_url(
        url,
        source_tier=source_tier,
        private_b0_domains=private_b0_domains,
    )
    domain = source_domain(url)
    transport = transport or _urllib_transport
    last_request = _DOMAIN_LAST_REQUEST.get(domain)
    current = clock()
    if last_request is not None and current - last_request < per_domain_interval:
        sleep(per_domain_interval - (current - last_request))
    _DOMAIN_LAST_REQUEST[domain] = clock()
    last_error = None
    for attempt in range(max_retries + 1):
        try:
            response = transport(
                url,
                method=request["method"],
                timeout=timeout_seconds,
                headers=request["headers"],
            )
            for redirect in response.get("redirects", []):
                validate_redirect(
                    url,
                    redirect,
                    source_tier=source_tier,
                    private_b0_domains=private_b0_domains,
                )
            validate_response_guard(
                status=response["status"],
                content_type=response["content_type"],
                content=response["body"],
                max_bytes=max_bytes,
                timeout_seconds=timeout_seconds,
            )
            return {
                "url": url,
                "domain": domain,
                "status": response["status"],
                "content_type": response["content_type"],
                "content": response["body"],
                "attempts": attempt + 1,
            }
        except IdentityWebEvidenceError:
            raise
        except (OSError, urllib.error.URLError) as exc:
            last_error = exc
            if attempt < max_retries:
                sleep(per_domain_interval)
    raise IdentityWebEvidenceError("PUBLIC_FETCH_FAILED") from last_error


def content_sha256(content):
    if not isinstance(content, bytes):
        raise IdentityWebEvidenceError("CONTENT_BYTES_REQUIRED")
    return hashlib.sha256(content).hexdigest()


def evidence_id(
    *,
    source_tier,
    source_kind,
    source_url,
    content_sha256_value,
    extractor_version=EXTRACTOR_VERSION,
):
    core = {
        "schema_version": SCHEMA_VERSION,
        "source_tier": source_tier,
        "source_kind": source_kind,
        "source_url": source_url,
        "content_sha256": _required_sha256(
            content_sha256_value,
            "CONTENT_SHA256",
        ),
        "extractor_version": extractor_version,
    }
    return "evidence_" + hashlib.sha256(
        canonical_bytes(core),
    ).hexdigest()[:24]


def normalize_claim(claim_type, raw_value):
    if claim_type not in CLAIM_TYPES:
        raise IdentityWebEvidenceError("CLAIM_TYPE_INVALID")
    text = _normalize_text(raw_value)
    if claim_type == "OFFICIAL_BIRTH_DATE":
        candidate = text.replace("/", "-").replace(".", "-")
        if not _DATE.fullmatch(candidate):
            raise IdentityWebEvidenceError("BIRTH_DATE_INVALID")
        text = candidate
    return text


def make_claim(
    *,
    claim_type,
    raw_value,
    source_locator,
    claim_confidence="MEDIUM",
    machine_extracted=True,
    human_verified=False,
):
    return {
        "claim_type": claim_type,
        "raw_value": _required_text(raw_value, "RAW_VALUE"),
        "normalized_value": normalize_claim(claim_type, raw_value),
        "source_locator": _required_text(
            source_locator,
            "SOURCE_LOCATOR",
        ),
        "claim_confidence": claim_confidence,
        "machine_extracted": machine_extracted,
        "human_verified": human_verified,
    }


def external_person_id_claim(
    *,
    namespace,
    identifier,
    source_semantics,
    source_locator,
):
    if source_semantics not in {"PLAYER_ENTITY", "PERSON_ENTITY"}:
        raise IdentityWebEvidenceError(
            "EXTERNAL_PERSON_SEMANTICS_REQUIRED",
        )
    namespace = _required_text(namespace, "EXTERNAL_ID_NAMESPACE")
    if not EXTERNAL_ID_NAMESPACE.fullmatch(namespace):
        raise IdentityWebEvidenceError("EXTERNAL_ID_NAMESPACE_INVALID")
    return make_claim(
        claim_type="OFFICIAL_SOURCE_DECLARED_PERSON_ID",
        raw_value=f"{namespace}:{_required_text(identifier, 'EXTERNAL_ID')}",
        source_locator=source_locator,
    )


def _validate_claim(claim):
    if not isinstance(claim, dict):
        raise IdentityWebEvidenceError("CLAIM_OBJECT_REQUIRED")
    required = {
        "claim_type",
        "raw_value",
        "normalized_value",
        "source_locator",
        "claim_confidence",
        "machine_extracted",
        "human_verified",
    }
    if set(claim) != required:
        raise IdentityWebEvidenceError("CLAIM_SCHEMA_INVALID")
    normalized = normalize_claim(claim["claim_type"], claim["raw_value"])
    if claim["normalized_value"] != normalized:
        raise IdentityWebEvidenceError("CLAIM_NORMALIZATION_MISMATCH")
    return claim


def _validate_evidence_binding(binding):
    if not isinstance(binding, dict):
        raise IdentityWebEvidenceError("EVIDENCE_BINDING_OBJECT_REQUIRED")
    required = {
        "record_key",
        "target_type",
        "target_id",
        "matched_claim_fields",
        "binding_rationale",
    }
    if set(binding) != required:
        raise IdentityWebEvidenceError("EVIDENCE_BINDING_SCHEMA_INVALID")
    _required_text(binding["record_key"], "RECORD_KEY")
    if binding["target_type"] not in {"EXISTING_UID", "NEW_GROUP"}:
        raise IdentityWebEvidenceError("EVIDENCE_TARGET_TYPE_INVALID")
    _required_text(binding["target_id"], "EVIDENCE_TARGET_ID")
    if not isinstance(binding["matched_claim_fields"], list):
        raise IdentityWebEvidenceError("MATCHED_CLAIM_FIELDS_REQUIRED")
    if binding["matched_claim_fields"] != sorted(
        set(binding["matched_claim_fields"]),
    ):
        raise IdentityWebEvidenceError("MATCHED_CLAIM_FIELDS_ORDER_INVALID")
    _required_text(binding["binding_rationale"], "BINDING_RATIONALE")
    return binding


def validate_evidence_item(item):
    if not isinstance(item, dict):
        raise IdentityWebEvidenceError("EVIDENCE_ITEM_OBJECT_REQUIRED")
    required = {
        "schema_version",
        "evidence_id",
        "source_tier",
        "source_kind",
        "source_url",
        "source_domain",
        "fetched_at",
        "http_status",
        "content_type",
        "content_sha256",
        "raw_snapshot_path",
        "published_at",
        "extractor_version",
        "claims",
        "extraction_warnings",
        "record_keys",
        "bindings",
        "b0_domain_approved",
    }
    if set(item) != required:
        raise IdentityWebEvidenceError("EVIDENCE_ITEM_SCHEMA_INVALID")
    if item["schema_version"] != SCHEMA_VERSION:
        raise IdentityWebEvidenceError("EVIDENCE_ITEM_VERSION_INVALID")
    expected_id = evidence_id(
        source_tier=item["source_tier"],
        source_kind=item["source_kind"],
        source_url=item["source_url"],
        content_sha256_value=item["content_sha256"],
        extractor_version=item["extractor_version"],
    )
    if item["evidence_id"] != expected_id:
        raise IdentityWebEvidenceError("EVIDENCE_ID_MISMATCH")
    if item["source_domain"] != source_domain(item["source_url"]):
        raise IdentityWebEvidenceError("EVIDENCE_DOMAIN_MISMATCH")
    if item["source_tier"] == "B0":
        if item["b0_domain_approved"] is not True:
            raise IdentityWebEvidenceError("B0_DOMAIN_NOT_APPROVED")
    else:
        validate_source_url(
            item["source_url"],
            source_tier=item["source_tier"],
        )
    if not isinstance(item["claims"], list):
        raise IdentityWebEvidenceError("CLAIMS_REQUIRED")
    for claim in item["claims"]:
        _validate_claim(claim)
    if item["extraction_warnings"] != sorted(set(item["extraction_warnings"])):
        raise IdentityWebEvidenceError("WARNING_ORDER_INVALID")
    if item["record_keys"] != sorted(set(item["record_keys"])):
        raise IdentityWebEvidenceError("RECORD_KEYS_ORDER_INVALID")
    if not isinstance(item["bindings"], list):
        raise IdentityWebEvidenceError("EVIDENCE_BINDINGS_REQUIRED")
    for binding in item["bindings"]:
        _validate_evidence_binding(binding)
    return item


def build_evidence_item(
    *,
    source_tier,
    source_kind,
    source_url,
    fetched_at,
    http_status,
    content_type,
    content,
    raw_snapshot_path,
    claims,
    record_keys,
    published_at=None,
    extraction_warnings=None,
    extractor_version=EXTRACTOR_VERSION,
    private_b0_domains=None,
    bindings=None,
    b0_domain_approved=False,
):
    validate_source_url(
        source_url,
        source_tier=source_tier,
        private_b0_domains=private_b0_domains,
    )
    digest = content_sha256(content)
    item_id = evidence_id(
        source_tier=source_tier,
        source_kind=source_kind,
        source_url=source_url,
        content_sha256_value=digest,
        extractor_version=extractor_version,
    )
    item = {
        "schema_version": SCHEMA_VERSION,
        "evidence_id": item_id,
        "source_tier": source_tier,
        "source_kind": _required_text(source_kind, "SOURCE_KIND"),
        "source_url": source_url,
        "source_domain": source_domain(source_url),
        "fetched_at": _required_text(fetched_at, "FETCHED_AT"),
        "http_status": http_status,
        "content_type": _required_text(content_type, "CONTENT_TYPE"),
        "content_sha256": digest,
        "raw_snapshot_path": _required_text(
            raw_snapshot_path,
            "RAW_SNAPSHOT_PATH",
        ),
        "published_at": published_at,
        "extractor_version": extractor_version,
        "claims": sorted(
            [_validate_claim(claim) for claim in claims],
            key=lambda item: (
                item["claim_type"],
                item["normalized_value"],
                item["source_locator"],
            ),
        ),
        "extraction_warnings": sorted(set(extraction_warnings or [])),
        "record_keys": sorted(set(record_keys)),
        "bindings": sorted(
            [_validate_evidence_binding(item) for item in bindings or []],
            key=lambda item: (
                item["record_key"],
                item["target_type"],
                item["target_id"],
            ),
        ),
        "b0_domain_approved": b0_domain_approved,
    }
    return validate_evidence_item(item)


class _ClaimHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.claims = []

    def handle_starttag(self, tag, attrs):
        data = dict(attrs)
        claim_type = data.get("data-claim-type")
        value = data.get("data-claim-value")
        if claim_type and value:
            self.claims.append(make_claim(
                claim_type=claim_type,
                raw_value=html.unescape(value),
                source_locator=data.get("data-claim-locator") or tag,
            ))


def extract_html_claims(content, *, source_locator="html"):
    if not isinstance(content, (bytes, str)):
        raise IdentityWebEvidenceError("HTML_CONTENT_REQUIRED")
    text = content.decode("utf-8") if isinstance(content, bytes) else content
    parser = _ClaimHTMLParser()
    parser.feed(text)
    return parser.claims


def extract_pdf_claims(content, *, text_extractor=None):
    if not isinstance(content, bytes):
        raise IdentityWebEvidenceError("PDF_BYTES_REQUIRED")
    if text_extractor is None:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise IdentityWebEvidenceError("PDF_EXTRACTOR_REQUIRED") from exc
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        text = text_extractor(content)
    claims = []
    for line_number, line in enumerate(text.splitlines(), 1):
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = key.strip()
        if key in CLAIM_TYPES:
            claims.append(make_claim(
                claim_type=key,
                raw_value=value,
                source_locator=f"pdf-line:{line_number}",
            ))
    return claims


def build_evidence_manifest(items):
    normalized = sorted(
        [validate_evidence_item(item) for item in items],
        key=lambda item: item["evidence_id"],
    )
    ids = [item["evidence_id"] for item in normalized]
    if len(ids) != len(set(ids)):
        raise IdentityWebEvidenceError("EVIDENCE_ID_DUPLICATE")
    core = {
        "schema_version": SCHEMA_VERSION,
        "extractor_version": EXTRACTOR_VERSION,
        "items": normalized,
    }
    return {
        **core,
        "web_evidence_manifest_sha256": hashlib.sha256(
            canonical_bytes(core),
        ).hexdigest(),
    }


def validate_evidence_manifest(value):
    if not isinstance(value, dict):
        raise IdentityWebEvidenceError("EVIDENCE_MANIFEST_OBJECT_REQUIRED")
    required = {
        "schema_version",
        "extractor_version",
        "items",
        "web_evidence_manifest_sha256",
    }
    if set(value) != required:
        raise IdentityWebEvidenceError("EVIDENCE_MANIFEST_SCHEMA_INVALID")
    core = {
        key: value[key]
        for key in value
        if key != "web_evidence_manifest_sha256"
    }
    if hashlib.sha256(canonical_bytes(core)).hexdigest() != value[
        "web_evidence_manifest_sha256"
    ]:
        raise IdentityWebEvidenceError("EVIDENCE_MANIFEST_HASH_MISMATCH")
    for item in value["items"]:
        validate_evidence_item(item)
    return value


def _claims_for_record(manifest, record_key):
    claims = []
    for item in manifest["items"]:
        if record_key not in item["record_keys"]:
            continue
        for claim in item["claims"]:
            claims.append({**claim, "evidence_id": item["evidence_id"]})
    return claims


def _claims_for_target(manifest, record_key, target_type, target_id):
    claims = []
    for item in manifest["items"]:
        if record_key not in item["record_keys"]:
            continue
        matching_binding = next((
            binding for binding in item["bindings"]
            if (
                binding["record_key"] == record_key
                and binding["target_type"] == target_type
                and binding["target_id"] == target_id
            )
        ), None)
        if matching_binding is None:
            continue
        for claim in item["claims"]:
            claims.append({
                **claim,
                "evidence_id": item["evidence_id"],
                "binding_rationale": matching_binding["binding_rationale"],
                "matched_claim_fields": list(
                    matching_binding["matched_claim_fields"],
                ),
            })
    return claims


def bind_evidence_to_target(
    *,
    record_key,
    target_type,
    target_id,
    matched_claim_fields,
    binding_rationale,
):
    return _validate_evidence_binding({
        "record_key": record_key,
        "target_type": target_type,
        "target_id": target_id,
        "matched_claim_fields": sorted(set(matched_claim_fields)),
        "binding_rationale": binding_rationale,
    })


def classify_candidate_evidence(
    record_key,
    *,
    target_type,
    target_id,
    manifest,
):
    validate_evidence_manifest(manifest)
    claims = _claims_for_target(
        manifest,
        record_key,
        target_type,
        target_id,
    )
    if not claims:
        return "W4_DISCOVERY_SUPPORT", ["NO_TARGET_BOUND_OFFICIAL_EVIDENCE"]
    evidence_ids = {claim["evidence_id"] for claim in claims}
    person_ids = _person_ids(claims)
    if len(person_ids) > 1:
        return "WX_CONFLICT", ["OFFICIAL_PERSON_ID_CONFLICT"]
    birth_dates = {
        claim["normalized_value"]
        for claim in claims
        if claim["claim_type"] == "OFFICIAL_BIRTH_DATE"
    }
    if len(birth_dates) > 1:
        return "WX_CONFLICT", ["OFFICIAL_BIRTH_DATE_CONFLICT"]
    if person_ids:
        return "W1_OFFICIAL_PERSON_ID_EXACT", []
    names = {
        claim["normalized_value"]
        for claim in claims
        if claim["claim_type"] == "OFFICIAL_PLAYER_NAME"
    }
    if names and birth_dates and len(evidence_ids) >= 2:
        return "W2_OFFICIAL_BIO_MULTI_SOURCE", []
    if names and any(
        claim["claim_type"] in {
            "OFFICIAL_TEAM",
            "OFFICIAL_SEASON",
            "OFFICIAL_REGISTRATION_UNIT",
        }
        for claim in claims
    ):
        return "W3_OFFICIAL_CONTINUITY", []
    return "W4_DISCOVERY_SUPPORT", []


def _person_ids(claims):
    return {
        claim["normalized_value"]
        for claim in claims
        if claim["claim_type"] == "OFFICIAL_SOURCE_DECLARED_PERSON_ID"
    }


def classify_record_evidence(record_key, manifest):
    validate_evidence_manifest(manifest)
    claims = _claims_for_record(manifest, record_key)
    evidence_ids = {claim["evidence_id"] for claim in claims}
    person_ids = _person_ids(claims)
    if len(person_ids) > 1:
        return "WX_CONFLICT", ["OFFICIAL_PERSON_ID_CONFLICT"]
    birth_dates = {
        claim["normalized_value"]
        for claim in claims
        if claim["claim_type"] == "OFFICIAL_BIRTH_DATE"
    }
    if len(birth_dates) > 1:
        return "WX_CONFLICT", ["OFFICIAL_BIRTH_DATE_CONFLICT"]
    if person_ids:
        return "W1_OFFICIAL_PERSON_ID_EXACT", []
    names = {
        claim["normalized_value"]
        for claim in claims
        if claim["claim_type"] == "OFFICIAL_PLAYER_NAME"
    }
    if names and birth_dates and len(evidence_ids) >= 2:
        return "W2_OFFICIAL_BIO_MULTI_SOURCE", []
    if names and any(
        claim["claim_type"] in {
            "OFFICIAL_TEAM",
            "OFFICIAL_SEASON",
            "OFFICIAL_REGISTRATION_UNIT",
        }
        for claim in claims
    ):
        return "W3_OFFICIAL_CONTINUITY", []
    return "W4_DISCOVERY_SUPPORT", []


def deterministic_group_id(
    member_review_ids,
    evidence_class,
    evidence_ids=None,
):
    core = {
        "member_review_ids": sorted(set(member_review_ids)),
        "evidence_class": evidence_class,
        "evidence_ids": sorted(set(evidence_ids or [])),
    }
    return "candidate_group_" + hashlib.sha256(
        canonical_bytes(core),
    ).hexdigest()[:24]


def build_conflicts(manifest):
    validate_evidence_manifest(manifest)
    conflicts = []
    record_keys = sorted({
        record_key
        for item in manifest["items"]
        for record_key in item["record_keys"]
    })
    for record_key in record_keys:
        evidence_class, reasons = classify_record_evidence(
            record_key,
            manifest,
        )
        if evidence_class == "WX_CONFLICT":
            conflicts.append({
                "record_key": record_key,
                "conflict_class": "WX_CONFLICT",
                "reasons": reasons,
            })
    return conflicts


def batch_predicate(proposal):
    if proposal.get("conflict_count", 0):
        return None
    if proposal.get("text_corruption_warning"):
        return None
    if (
        proposal.get("proposal_type") == "EXISTING_IDENTITY_CANDIDATE"
        and proposal.get("evidence_class") == "W1_OFFICIAL_PERSON_ID_EXACT"
    ):
        return "BATCH_EXISTING_W1"
    if (
        proposal.get("proposal_type") == "EXISTING_IDENTITY_CANDIDATE"
        and proposal.get("evidence_class") == "W2_OFFICIAL_BIO_MULTI_SOURCE"
    ):
        return "BATCH_EXISTING_W2"
    if (
        proposal.get("proposal_type") == "NO_SAFE_CANDIDATE"
        and proposal.get("search_status")
        == "OFFICIAL_SEARCH_EXHAUSTED_NO_SAFE_MATCH"
    ):
        return "BATCH_NO_SAFE_SEARCH_EXHAUSTED"
    return None


def build_collector_search_manifest(
    *,
    record_key,
    applicable_required_collectors,
    completed_collectors,
    failed_collectors,
    attempted_official_urls,
    collector_version,
    text_warning,
    identity_conflict,
    surviving_candidate_count,
):
    core = {
        "schema_version": SCHEMA_VERSION,
        "record_key": record_key,
        "applicable_required_collectors": sorted(
            set(applicable_required_collectors),
        ),
        "completed_collectors": sorted(set(completed_collectors)),
        "failed_collectors": sorted(set(failed_collectors)),
        "attempted_official_urls": sorted(set(attempted_official_urls)),
        "collector_version": collector_version,
        "text_warning": bool(text_warning),
        "identity_conflict": bool(identity_conflict),
        "surviving_candidate_count": surviving_candidate_count,
    }
    return {
        **core,
        "collector_search_manifest_sha256": hashlib.sha256(
            canonical_bytes(core),
        ).hexdigest(),
    }


def validate_collector_search_manifest(value):
    if not isinstance(value, dict):
        raise IdentityWebEvidenceError("SEARCH_MANIFEST_OBJECT_REQUIRED")
    required = {
        "schema_version",
        "record_key",
        "applicable_required_collectors",
        "completed_collectors",
        "failed_collectors",
        "attempted_official_urls",
        "collector_version",
        "text_warning",
        "identity_conflict",
        "surviving_candidate_count",
        "collector_search_manifest_sha256",
    }
    if set(value) != required:
        raise IdentityWebEvidenceError("SEARCH_MANIFEST_SCHEMA_INVALID")
    core = {
        key: value[key]
        for key in value
        if key != "collector_search_manifest_sha256"
    }
    if hashlib.sha256(canonical_bytes(core)).hexdigest() != value[
        "collector_search_manifest_sha256"
    ]:
        raise IdentityWebEvidenceError("SEARCH_MANIFEST_HASH_MISMATCH")
    return value


def search_exhaustion_status(search_manifest):
    validate_collector_search_manifest(search_manifest)
    eligible = (
        search_manifest["applicable_required_collectors"]
        and set(search_manifest["applicable_required_collectors"])
        == set(search_manifest["completed_collectors"])
        and not search_manifest["failed_collectors"]
        and not search_manifest["text_warning"]
        and not search_manifest["identity_conflict"]
        and search_manifest["surviving_candidate_count"] == 0
    )
    return (
        "OFFICIAL_SEARCH_EXHAUSTED_NO_SAFE_MATCH"
        if eligible
        else "SEARCH_INCOMPLETE_OR_CONFLICT"
    )


def enrich_candidates(
    candidates,
    *,
    evidence_manifest,
    search_statuses=None,
):
    validate_evidence_manifest(evidence_manifest)
    search_statuses = search_statuses or {}
    conflicts = {
        item["record_key"] for item in build_conflicts(evidence_manifest)
    }
    enriched = []
    for proposal in candidates:
        record_key = proposal["record_key"]
        search_manifest = search_statuses.get(record_key)
        search_status = (
            search_exhaustion_status(search_manifest)
            if isinstance(search_manifest, dict)
            else "SEARCH_NOT_ATTEMPTED"
        )
        if proposal.get("candidate_player_uid"):
            evidence_class, reasons = classify_candidate_evidence(
                record_key,
                target_type="EXISTING_UID",
                target_id=proposal["candidate_player_uid"],
                manifest=evidence_manifest,
            )
        elif proposal.get("candidate_group_id"):
            evidence_class, reasons = classify_candidate_evidence(
                record_key,
                target_type="NEW_GROUP",
                target_id=proposal["candidate_group_id"],
                manifest=evidence_manifest,
            )
        else:
            evidence_class, reasons = "W4_DISCOVERY_SUPPORT", [
                "NO_EXPLICIT_CANDIDATE_TARGET",
            ]
        item = {
            **proposal,
            "evidence_class": evidence_class,
            "conflict_count": int(evidence_class == "WX_CONFLICT"),
            "conflict_reasons": reasons,
            "search_status": search_status,
            "text_corruption_warning": any(
                _SUSPICIOUS_TEXT.search(claim["raw_value"])
                for claim in _claims_for_record(evidence_manifest, record_key)
            ),
            "batch_predicate": None,
        }
        if proposal["proposal_type"] == "NO_SAFE_CANDIDATE":
            if item["search_status"] == (
                "OFFICIAL_SEARCH_EXHAUSTED_NO_SAFE_MATCH"
            ):
                item["proposal_type"] = "NO_SAFE_CANDIDATE"
                item["machine_reason"] = item["search_status"]
            else:
                item["proposal_type"] = "NO_SAFE_CANDIDATE"
                item["proposed_relation"] = "NO_SAFE_CANDIDATE"
                item["machine_reason"] = "no_target_bound_safe_candidate"
        item["batch_predicate"] = batch_predicate(item)
        enriched.append(item)
    return sorted(
        enriched,
        key=lambda item: (
            item["record_key"],
            item["proposal_type"],
            item.get("candidate_player_uid") or "",
            item.get("candidate_group_id") or "",
        ),
    )


def build_batch_plan(
    packet,
    evidence_manifest,
    *,
    batch_size=200,
    search_statuses=None,
):
    validate_review_packet(packet)
    validate_evidence_manifest(evidence_manifest)
    search_statuses = search_statuses or {}
    eligible = []
    for proposal in packet["reviews"]:
        claims = _claims_for_record(evidence_manifest, proposal["record_key"])
        evidence_class, reasons = classify_record_evidence(
            proposal["record_key"],
            evidence_manifest,
        )
        search_manifest = search_statuses.get(proposal["record_key"])
        search_status = (
            search_exhaustion_status(search_manifest)
            if isinstance(search_manifest, dict)
            else "SEARCH_NOT_ATTEMPTED"
        )
        item = {
            **proposal,
            "evidence_class": evidence_class,
            "conflict_count": int(evidence_class == "WX_CONFLICT"),
            "text_corruption_warning": any(
                _SUSPICIOUS_TEXT.search(claim["raw_value"])
                for claim in claims
            ),
            "search_status": search_status,
        }
        pred = batch_predicate(item)
        if pred is not None:
            eligible.append({**item, "batch_predicate": pred})
    eligible.sort(key=lambda item: (
        item["batch_predicate"],
        item["record_key"],
        item["review_id"],
    ))
    batches = []
    by_predicate = defaultdict(list)
    for item in eligible:
        by_predicate[item["batch_predicate"]].append(item)
    for predicate in sorted(by_predicate):
        predicate_rows = sorted(
            by_predicate[predicate],
            key=lambda item: (item["record_key"], item["review_id"]),
        )
        for start in range(0, len(predicate_rows), batch_size):
            members = predicate_rows[start:start + batch_size]
            batches.append({
                "batch_id": "batch_" + hashlib.sha256(
                    canonical_bytes([
                        item["review_id"] for item in members
                    ]),
                ).hexdigest()[:24],
                "batch_predicate": predicate,
                "member_review_ids": [item["review_id"] for item in members],
                "member_count": len(members),
            })
    return {
        "schema_version": SCHEMA_VERSION,
        "authority": "NON_CANONICAL_HUMAN_WORKFLOW_PROJECTION",
        "review_packet_sha256": packet["review_packet_sha256"],
        "web_evidence_manifest_sha256": evidence_manifest[
            "web_evidence_manifest_sha256"
        ],
        "batch_size": batch_size,
        "batches": batches,
    }


def build_groups(packet, evidence_manifest):
    validate_review_packet(packet)
    validate_evidence_manifest(evidence_manifest)
    grouped = defaultdict(list)
    for proposal in packet["reviews"]:
        group_id = proposal.get("candidate_group_id")
        if not group_id:
            continue
        evidence_class, reasons = classify_candidate_evidence(
            proposal["record_key"],
            target_type="NEW_GROUP",
            target_id=group_id,
            manifest=evidence_manifest,
        )
        if (
            proposal["proposal_type"] == "NEW_IDENTITY_CANDIDATE"
            and evidence_class in {
                "W1_OFFICIAL_PERSON_ID_EXACT",
                "W2_OFFICIAL_BIO_MULTI_SOURCE",
            }
            and not reasons
        ):
            grouped[group_id].append(
                proposal["review_id"],
            )
    return [
        {
            "candidate_group_id": group_id,
            "member_review_ids": sorted(member_ids),
            "member_count": len(member_ids),
            "evidence_class": classify_candidate_evidence(
                next(
                    item["record_key"]
                    for item in packet["reviews"]
                    if item["candidate_group_id"] == group_id
                ),
                target_type="NEW_GROUP",
                target_id=group_id,
                manifest=evidence_manifest,
            )[0],
            "conflict_count": 0,
        }
        for group_id, member_ids in sorted(grouped.items())
    ]


def plan_sha256(value):
    return hashlib.sha256(canonical_bytes(value)).hexdigest()


def _validate_authority_event_common(event, packet, evidence_manifest):
    if packet["review_packet_sha256"] != event["review_packet_sha256"]:
        raise IdentityWebEvidenceError("REVIEW_PACKET_HASH_CHANGED")
    if evidence_manifest["web_evidence_manifest_sha256"] != event[
        "web_evidence_manifest_sha256"
    ]:
        raise IdentityWebEvidenceError("EVIDENCE_MANIFEST_HASH_CHANGED")
    if event["human_decision"] != "APPROVE":
        raise IdentityWebEvidenceError("HUMAN_APPROVAL_REQUIRED")
    _required_text(event["human_note"], "HUMAN_NOTE")
    _required_text(event["reviewed_at"], "REVIEWED_AT")


def expand_group_event(packet, evidence_manifest, groups, event):
    validate_review_packet(packet)
    validate_evidence_manifest(evidence_manifest)
    _validate_authority_event_common(event, packet, evidence_manifest)
    group = next(
        (
            item for item in groups
            if item["candidate_group_id"] == event["candidate_group_id"]
        ),
        None,
    )
    if group is None:
        raise IdentityWebEvidenceError("GROUP_UNKNOWN")
    if group["member_review_ids"] != sorted(set(event["member_review_ids"])):
        raise IdentityWebEvidenceError("GROUP_MEMBERSHIP_MISMATCH")
    if group["evidence_class"] not in {
        "W1_OFFICIAL_PERSON_ID_EXACT",
        "W2_OFFICIAL_BIO_MULTI_SOURCE",
    }:
        raise IdentityWebEvidenceError("GROUP_NOT_ELIGIBLE")
    validate_player_uid(event["approved_player_uid"])
    _required_text(event["approved_canonical_name"], "APPROVED_CANONICAL_NAME")
    _required_text(event["allocation_attestation"], "ALLOCATION_ATTESTATION")
    review_by_id = {item["review_id"]: item for item in packet["reviews"]}
    decisions = []
    for review_id in group["member_review_ids"]:
        item = review_by_id[review_id]
        decisions.append({
            **item,
            "human_decision": "APPROVE",
            "human_note": (
                f"{event['human_note']} "
                f"[authority_event_id={event['event_id']}]"
            ),
            "approved_player_uid": event["approved_player_uid"],
            "approved_canonical_name": event["approved_canonical_name"],
            "source_exception_reason": None,
            "reviewed_at": event["reviewed_at"],
            "authority_event_id": event["event_id"],
        })
    return decisions


def expand_batch_event(packet, evidence_manifest, batches, event):
    validate_review_packet(packet)
    validate_evidence_manifest(evidence_manifest)
    _validate_authority_event_common(event, packet, evidence_manifest)
    batch = next(
        (
            item for item in batches
            if item["batch_id"] == event["batch_id"]
        ),
        None,
    )
    if batch is None:
        raise IdentityWebEvidenceError("BATCH_UNKNOWN")
    if batch["member_review_ids"] != sorted(set(event["member_review_ids"])):
        raise IdentityWebEvidenceError("BATCH_MEMBERSHIP_MISMATCH")
    if batch["batch_predicate"] not in BATCH_PREDICATES:
        raise IdentityWebEvidenceError("BATCH_PREDICATE_INVALID")
    review_by_id = {item["review_id"]: item for item in packet["reviews"]}
    decisions = []
    for review_id in batch["member_review_ids"]:
        item = review_by_id[review_id]
        if item["proposal_type"] == "NEW_IDENTITY_CANDIDATE":
            raise IdentityWebEvidenceError("NEW_IDENTITY_BATCH_FORBIDDEN")
        decisions.append({
            **item,
            "human_decision": "APPROVE",
            "human_note": (
                f"{event['human_note']} "
                f"[authority_event_id={event['event_id']}]"
            ),
            "approved_player_uid": None,
            "approved_canonical_name": None,
            "source_exception_reason": None,
            "reviewed_at": event["reviewed_at"],
            "authority_event_id": event["event_id"],
        })
    return decisions


def _write_sheet(sheet, headers, rows, machine_fields=()):
    sheet.append(headers)
    for row in rows:
        sheet.append([
            (
                json.dumps(
                    row.get(header),
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                if isinstance(row.get(header), (list, dict))
                else row.get(header)
            )
            for header in headers
        ])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            header = headers[cell.column - 1]
            if header in machine_fields:
                cell.protection = Protection(locked=True)
            else:
                cell.protection = Protection(locked=False)
    sheet.protection.sheet = True


def export_review_workbook(
    path,
    *,
    packet,
    evidence_manifest,
    groups,
    batches,
):
    validate_review_packet(packet)
    validate_evidence_manifest(evidence_manifest)
    workbook = Workbook()
    workbook.remove(workbook.active)
    review = workbook.create_sheet("Review")
    _write_sheet(
        review,
        list(REVIEW_FIELDS),
        packet["reviews"],
        machine_fields=MACHINE_REVIEW_FIELDS,
    )
    groups_sheet = workbook.create_sheet("Groups")
    groups_plan_sha = plan_sha256(groups)
    _write_sheet(
        groups_sheet,
        [
            "candidate_group_id",
            "member_review_ids",
            "member_count",
            "evidence_class",
            "conflict_count",
            "human_decision",
            "human_note",
            "approved_player_uid",
            "approved_canonical_name",
            "allocation_attestation",
            "reviewed_at",
            "plan_sha256",
        ],
        [
            {
                **group,
                "plan_sha256": groups_plan_sha,
                "member_review_ids": json.dumps(
                    group["member_review_ids"],
                    separators=(",", ":"),
                ),
            }
            for group in groups
        ],
        machine_fields={
            "candidate_group_id",
            "member_review_ids",
            "member_count",
            "evidence_class",
            "conflict_count",
            "plan_sha256",
        },
    )
    batches_sheet = workbook.create_sheet("Batches")
    batches_plan_sha = plan_sha256(batches)
    _write_sheet(
        batches_sheet,
        [
            "batch_id",
            "batch_predicate",
            "member_review_ids",
            "member_count",
            "human_decision",
            "human_note",
            "reviewed_at",
            "plan_sha256",
        ],
        [
            {
                **batch,
                "plan_sha256": batches_plan_sha,
                "member_review_ids": json.dumps(
                    batch["member_review_ids"],
                    separators=(",", ":"),
                ),
            }
            for batch in batches
        ],
        machine_fields={
            "batch_id",
            "batch_predicate",
            "member_review_ids",
            "member_count",
            "plan_sha256",
        },
    )
    evidence_sheet = workbook.create_sheet("Evidence_Index")
    evidence_rows = [
        {
            "evidence_id": item["evidence_id"],
            "source_tier": item["source_tier"],
            "source_url": item["source_url"],
            "source_kind": item["source_kind"],
            "content_sha256": item["content_sha256"],
        }
        for item in evidence_manifest["items"]
    ]
    _write_sheet(
        evidence_sheet,
        [
            "evidence_id",
            "source_tier",
            "source_url",
            "source_kind",
            "content_sha256",
        ],
        evidence_rows,
        machine_fields={
            "evidence_id",
            "source_tier",
            "source_url",
            "source_kind",
            "content_sha256",
        },
    )
    instructions = workbook.create_sheet("Instructions")
    for row in INSTRUCTIONS_ROWS:
        instructions.append(list(row))
    instructions.protection.sheet = True
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _sheet_table(path, sheet_name):
    workbook = load_workbook(path, read_only=False)
    sheet = workbook[sheet_name]
    values = list(sheet.values)
    headers = list(values[0])
    rows = [
        dict(zip(headers, row))
        for row in values[1:]
        if any(value is not None for value in row)
    ]
    return headers, rows


def _sheet_rows(path, sheet_name):
    return _sheet_table(path, sheet_name)[1]


def _verify_workbook_tables(path, *, packet, evidence_manifest, groups, batches):
    workbook = load_workbook(path, read_only=False)
    if tuple(workbook.sheetnames) != WORKBOOK_SHEETS:
        raise IdentityWebEvidenceError("WORKBOOK_SHEET_SET_INVALID")
    review_headers, review_rows = _sheet_table(path, "Review")
    if review_headers != list(REVIEW_FIELDS):
        raise IdentityWebEvidenceError("WORKBOOK_REVIEW_HEADERS_INVALID")
    group_headers, group_rows = _sheet_table(path, "Groups")
    expected_group_headers = [
        "candidate_group_id",
        "member_review_ids",
        "member_count",
        "evidence_class",
        "conflict_count",
        "human_decision",
        "human_note",
        "approved_player_uid",
        "approved_canonical_name",
        "allocation_attestation",
        "reviewed_at",
        "plan_sha256",
    ]
    if group_headers != expected_group_headers:
        raise IdentityWebEvidenceError("WORKBOOK_GROUP_HEADERS_INVALID")
    batch_headers, batch_rows = _sheet_table(path, "Batches")
    expected_batch_headers = [
        "batch_id",
        "batch_predicate",
        "member_review_ids",
        "member_count",
        "human_decision",
        "human_note",
        "reviewed_at",
        "plan_sha256",
    ]
    if batch_headers != expected_batch_headers:
        raise IdentityWebEvidenceError("WORKBOOK_BATCH_HEADERS_INVALID")
    evidence_headers, evidence_rows = _sheet_table(path, "Evidence_Index")
    expected_evidence_headers = [
        "evidence_id",
        "source_tier",
        "source_url",
        "source_kind",
        "content_sha256",
    ]
    if evidence_headers != expected_evidence_headers:
        raise IdentityWebEvidenceError("WORKBOOK_EVIDENCE_HEADERS_INVALID")
    instructions_headers, instructions_rows = _sheet_table(
        path,
        "Instructions",
    )
    if instructions_headers != ["Authority boundary"]:
        raise IdentityWebEvidenceError("WORKBOOK_INSTRUCTIONS_INVALID")
    actual_instructions = tuple(
        (row.get("Authority boundary"),)
        for row in instructions_rows
    )
    if actual_instructions != INSTRUCTIONS_ROWS[1:]:
        raise IdentityWebEvidenceError("WORKBOOK_INSTRUCTIONS_INVALID")

    group_by_id = {item["candidate_group_id"]: item for item in groups}
    expected_groups_plan_sha = plan_sha256(groups)
    for row in group_rows:
        expected = group_by_id.get(row["candidate_group_id"])
        if expected is None:
            raise IdentityWebEvidenceError("GROUP_UNKNOWN")
        for field in (
            "member_review_ids",
            "member_count",
            "evidence_class",
            "conflict_count",
        ):
            actual = row[field]
            if field == "member_review_ids" and isinstance(actual, str):
                actual = json.loads(actual)
            if actual != expected[field]:
                raise IdentityWebEvidenceError(
                    f"WORKBOOK_MACHINE_FIELD_TAMPER:Groups.{field}",
                )
        if row["plan_sha256"] != expected_groups_plan_sha:
            raise IdentityWebEvidenceError(
                "WORKBOOK_MACHINE_FIELD_TAMPER:Groups.plan_sha256",
            )

    batch_by_id = {item["batch_id"]: item for item in batches}
    expected_batches_plan_sha = plan_sha256(batches)
    for row in batch_rows:
        expected = batch_by_id.get(row["batch_id"])
        if expected is None:
            raise IdentityWebEvidenceError("BATCH_UNKNOWN")
        for field in ("batch_predicate", "member_review_ids", "member_count"):
            actual = row[field]
            if field == "member_review_ids" and isinstance(actual, str):
                actual = json.loads(actual)
            if actual != expected[field]:
                raise IdentityWebEvidenceError(
                    f"WORKBOOK_MACHINE_FIELD_TAMPER:Batches.{field}",
                )
        if row["plan_sha256"] != expected_batches_plan_sha:
            raise IdentityWebEvidenceError(
                "WORKBOOK_MACHINE_FIELD_TAMPER:Batches.plan_sha256",
            )

    expected_evidence = {
        item["evidence_id"]: {
            "source_tier": item["source_tier"],
            "source_url": item["source_url"],
            "source_kind": item["source_kind"],
            "content_sha256": item["content_sha256"],
        }
        for item in evidence_manifest["items"]
    }
    if len(evidence_rows) != len(expected_evidence):
        raise IdentityWebEvidenceError("WORKBOOK_EVIDENCE_ROWS_MISMATCH")
    for row in evidence_rows:
        expected = expected_evidence.get(row["evidence_id"])
        if expected is None:
            raise IdentityWebEvidenceError("WORKBOOK_EVIDENCE_UNKNOWN")
        for field, value in expected.items():
            if row[field] != value:
                raise IdentityWebEvidenceError(
                    f"WORKBOOK_MACHINE_FIELD_TAMPER:Evidence_Index.{field}",
                )


def import_review_workbook(
    path,
    *,
    packet,
    evidence_manifest,
    groups,
    batches,
):
    validate_review_packet(packet)
    validate_evidence_manifest(evidence_manifest)
    _verify_workbook_tables(
        path,
        packet=packet,
        evidence_manifest=evidence_manifest,
        groups=groups,
        batches=batches,
    )
    review_rows = _sheet_rows(path, "Review")
    review_by_id = {item["review_id"]: item for item in packet["reviews"]}
    if {row["review_id"] for row in review_rows} != set(review_by_id):
        raise IdentityWebEvidenceError("WORKBOOK_REVIEW_ROWS_MISMATCH")
    decisions = {}
    for row in review_rows:
        authority = review_by_id[row["review_id"]]
        for field in MACHINE_REVIEW_FIELDS:
            actual = row.get(field)
            expected = authority[field]
            if field == "evidence_refs" and isinstance(actual, str):
                actual = json.loads(actual or "[]")
            if actual != expected:
                raise IdentityWebEvidenceError(
                    f"WORKBOOK_MACHINE_FIELD_TAMPER:{field}",
                )
        if row.get("human_decision"):
            decisions[row["review_id"]] = {
                **authority,
                "human_decision": row["human_decision"],
                "human_note": row.get("human_note"),
                "approved_player_uid": row.get("approved_player_uid"),
                "approved_canonical_name": row.get("approved_canonical_name"),
                "source_exception_reason": row.get("source_exception_reason"),
                "reviewed_at": row.get("reviewed_at"),
            }

    group_rows = _sheet_rows(path, "Groups")
    group_by_id = {item["candidate_group_id"]: item for item in groups}
    for row in group_rows:
        if not row.get("human_decision"):
            continue
        group_id = row["candidate_group_id"]
        expected = group_by_id.get(group_id)
        if expected is None:
            raise IdentityWebEvidenceError("GROUP_UNKNOWN")
        members = json.loads(row["member_review_ids"] or "[]")
        if members != expected["member_review_ids"]:
            raise IdentityWebEvidenceError("GROUP_MEMBERSHIP_MISMATCH")
        event = {
            "event_id": f"group-event-{group_id}",
            "candidate_group_id": group_id,
            "review_packet_sha256": packet["review_packet_sha256"],
            "web_evidence_manifest_sha256": evidence_manifest[
                "web_evidence_manifest_sha256"
            ],
            "member_review_ids": members,
            "human_decision": row["human_decision"],
            "human_note": row["human_note"],
            "approved_player_uid": row.get("approved_player_uid"),
            "approved_canonical_name": row.get("approved_canonical_name"),
            "allocation_attestation": row.get("allocation_attestation"),
            "reviewed_at": row["reviewed_at"],
        }
        for item in expand_group_event(
            packet,
            evidence_manifest,
            groups,
            event,
        ):
            decisions[item["review_id"]] = item

    batch_rows = _sheet_rows(path, "Batches")
    batch_by_id = {item["batch_id"]: item for item in batches}
    for row in batch_rows:
        if not row.get("human_decision"):
            continue
        batch_id = row["batch_id"]
        expected = batch_by_id.get(batch_id)
        if expected is None:
            raise IdentityWebEvidenceError("BATCH_UNKNOWN")
        members = json.loads(row["member_review_ids"] or "[]")
        if members != expected["member_review_ids"]:
            raise IdentityWebEvidenceError("BATCH_MEMBERSHIP_MISMATCH")
        event = {
            "event_id": f"batch-event-{batch_id}",
            "batch_id": batch_id,
            "review_packet_sha256": packet["review_packet_sha256"],
            "web_evidence_manifest_sha256": evidence_manifest[
                "web_evidence_manifest_sha256"
            ],
            "member_review_ids": members,
            "batch_predicate": row["batch_predicate"],
            "human_decision": row["human_decision"],
            "human_note": row["human_note"],
            "reviewed_at": row["reviewed_at"],
        }
        for item in expand_batch_event(
            packet,
            evidence_manifest,
            batches,
            event,
        ):
            if item["review_id"] in decisions:
                raise IdentityWebEvidenceError("WORKBOOK_DECISION_CONFLICT")
            decisions[item["review_id"]] = item

    output = io.StringIO(newline="")
    writer = csv.DictWriter(
        output,
        fieldnames=list(REVIEW_FIELDS),
        lineterminator="\n",
    )
    writer.writeheader()
    for review_id in [item["review_id"] for item in packet["reviews"]]:
        item = decisions.get(review_id, review_by_id[review_id])
        writer.writerow({
            field: (
                json.dumps(
                    item[field],
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                if field == "evidence_refs"
                else ("" if item.get(field) is None else item.get(field))
            )
            for field in REVIEW_FIELDS
        })
    return output.getvalue().encode("utf-8")
