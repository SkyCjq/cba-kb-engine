"""Read-only MASTER validation and deterministic AI exports."""
import csv
import io
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from .common import atomic, digest, save

MASTER_ID = '1Nb4-4rrySjKW7GSA_PLh1SCkPgW9kJtc'
HEADERS = ['record_key','season','club_id','club_official','sequence','player',
           'registration_stage','registration_method','contract_category',
           'contract_term_official','former_club','notice_deadline',
           'registration_status','remarks','source_file_id','source_url',
           'source_page','source_type','extraction_method','verification_level']


def inspect(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != ['MASTER']:
            raise ValueError('Expected only MASTER tab')
        values = list(book['MASTER'].values)
    finally:
        book.close()
    if list(values[0]) != HEADERS:
        raise ValueError('MASTER schema changed; review required')
    rows = []
    for index, values_row in enumerate(values[1:], 2):
        if not any(v is not None for v in values_row):
            continue
        if len(values_row) != len(HEADERS):
            raise ValueError(f'Invalid width at row {index}')
        if any(v is not None and not isinstance(v, (str, int, float, bool)) for v in values_row):
            raise ValueError(f'Unexpected typed date/object at row {index}')
        row = dict(zip(HEADERS, values_row))
        for key in ['record_key','season','club_id','player','source_url','verification_level']:
            if not row[key]:
                raise ValueError(f'Missing {key} at row {index}')
        if row['record_key'] != '|'.join(str(row[k]) for k in ['season','club_id','player']):
            raise ValueError(f'Record key mismatch at row {index}')
        rows.append(row)
    if not rows or len({r['record_key'] for r in rows}) != len(rows):
        raise ValueError('Empty or duplicate records')
    normalized = json.dumps(rows, ensure_ascii=False, separators=(',',':')).encode()
    summary = {'rows':len(rows), 'columns':len(HEADERS), 'unique_keys':len(rows),
               'seasons':dict(sorted(Counter(r['season'] for r in rows).items())),
               'verification_levels':dict(Counter(r['verification_level'] for r in rows)),
               'sha256':digest(Path(path).read_bytes()), 'semantic_sha256':digest(normalized),
               'master_file_id':MASTER_ID}
    return rows, summary


def build(master, output, release_id, commit, generated_at):
    output = Path(output)
    if output.exists() and any(output.iterdir()):
        raise ValueError('Build directory must be new; preserve frozen candidates')
    rows, summary = inspect(master)
    meta = dict(release_id=release_id, commit_sha=commit, generated_at=generated_at,
                source_master_file_id=MASTER_ID, source_master_sha256=summary['sha256'], generated=True)
    output.mkdir(parents=True, exist_ok=True)
    save(output/'provenance.json', meta)
    save(output/'validation.json', summary)
    atomic(output/'MASTER.jsonl', ('\n'.join(json.dumps(r,ensure_ascii=False) for r in rows)+'\n').encode())
    stream = io.StringIO(newline='')
    writer = csv.DictWriter(stream,fieldnames=HEADERS,lineterminator='\n')
    writer.writeheader(); writer.writerows(rows)
    atomic(output/'MASTER.csv', stream.getvalue().encode('utf-8-sig'))
    preamble = ('GENERATED / READ-ONLY\n'+json.dumps(meta,ensure_ascii=False)+'\n\n'
                'auto_validated 为机器核验，不等于官方或人工核验。空字段表示无已登记值，不是零。\n\n')
    index = ['# CBA-KB AI INDEX',preamble,
             f'MASTER: https://drive.google.com/file/d/{MASTER_ID}/view',
             f'记录数：{summary["rows"]}；这是注册记录数，不是独立球员人数。',
             '精确查询读取 MASTER.xlsx 或同版本 MASTER.jsonl；CSV 将空值表现为空单元格。',
             '读取 release_status.json：发布未完成时使用 previous_snapshot，不混用版本。','']
    for season, count in summary['seasons'].items():
        text = [f'# CBA 注册 {season}',preamble,f'注册记录数：{count}','']
        for row in rows:
            if row['season'] == season:
                text += [f'## {row["club_official"]} / {row["player"]}',
                         *[f'- {k}: {row[k] if row[k] is not None else "（空）"}' for k in HEADERS], '']
        atomic(output/f'CBA_注册_{season}.md', '\n'.join(text).encode())
        index.append(f'- CBA_注册_{season}.md：{count} 条')
    atomic(output/'INDEX.md', '\n'.join(index).encode())
    return summary
