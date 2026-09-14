"""Adapter for the CBA foreign player registration workbook."""
from pathlib import Path
import re
from openpyxl import load_workbook
from . import foreign

METHOD = 'xlsx merge-header segmentation + season-aware cancellation parsing'


def rows(path, sheet=None):
    book = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        name = sheet or book.sheetnames[0]
        if book.sheetnames != [name]:
            raise ValueError('Foreign source must stay a single-sheet file')
        cells = list(book[name].iter_rows())
        table = [[cell.value for cell in row] for row in cells]
        index = foreign.header_index(table)
        jersey = foreign.columns(table, index)['球衣号码']
        for number in range(index + 2, len(cells)):
            cell = cells[number][jersey]
            value = cell.value
            # A numeric zero displayed with the format 00 is a text code 00.
            # Keep General-format zero as 0, and do not coerce source strings.
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                if re.fullmatch(r'0+', cell.number_format or '') and value == int(value):
                    table[number][jersey] = str(int(value)).zfill(len(cell.number_format))
        return name, table
    finally:
        book.close()


def extract(path, season, source, clubs=None, strict=True, sheet=None):
    page, table = rows(path, sheet)
    title = str(table[0][0]) if table and table[0] and table[0][0] else None
    source = {**source, 'method': source.get('method') or METHOD}
    return foreign.build(table, season, source, clubs=clubs, strict=strict,
                         title=title, source_page=page)
