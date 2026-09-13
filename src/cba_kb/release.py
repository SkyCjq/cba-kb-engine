"""Frozen release plans, optimistic checks, journaling and rollback.

Cross-client exclusion still requires a human-maintained single-writer window.
Native Docs require explicit managed_doc mode; Sheets byte writes are rejected.
"""
from pathlib import Path
import json
import re
import subprocess
import time
from .common import atomic, child, digest, lock, read, save
from .transport import stage
from .current_state import clean

FOLDER = 'application/vnd.google-apps.folder'
DOC = 'application/vnd.google-apps.document'

RELEASE_STATE_SEQUENCE = (
    'PROJECTED',
    'RESERVED',
    'REPROJECTED',
    'FROZEN_PLAN',
    'PREPARED',
    'ARCHIVING',
    'ARCHIVE_COMPLETE',
    'PUBLISHING',
    'VERIFYING',
    'COMPLETE',
)
RELEASE_STATES = frozenset(RELEASE_STATE_SEQUENCE)
SHA_ROLES = (
    'baseline_development_sha',
    'product_candidate_sha',
    'release_execution_sha',
    'reviewed_release_pr_head_sha',
    'reviewed_ci_head_sha',
    'release_merge_sha',
    'production_execution_sha',
    'release_critical_tree_attestation',
)
RETRYABLE = 'RETRYABLE'
NON_RETRYABLE = 'NON_RETRYABLE'
RETRYABLE_STATUS = frozenset({408, 429})
NON_RETRYABLE_TERMS = (
    'auth', 'permission', 'forbidden', 'unauthorized', 'schema', 'authority',
    'hash', 'mime', 'semantic drift', 'semantic_drift',
)
RETRYABLE_TERMS = (
    'timeout', 'timed out', 'connection reset', 'broken pipe',
    'incomplete read', 'remote end closed',
)


class ReleaseContractError(RuntimeError):
    pass


class PreMutationAbort(ReleaseContractError):
    pass


class PartialMutationError(ReleaseContractError):
    pass


def _sha(value, role):
    if not isinstance(value, str) or not re.fullmatch('[0-9a-f]{40}', value):
        raise ReleaseContractError(f'{role.upper()}_INVALID')
    return value


def validate_state_transition(current, target):
    if current not in RELEASE_STATES or target not in RELEASE_STATES:
        raise ReleaseContractError('RELEASE_STATE_INVALID')
    current_index = RELEASE_STATE_SEQUENCE.index(current)
    target_index = RELEASE_STATE_SEQUENCE.index(target)
    if target_index < current_index:
        raise ReleaseContractError('RELEASE_STATE_REGRESSION')
    if target_index > current_index + 1:
        raise ReleaseContractError('RELEASE_STATE_SKIP')
    return True


def provenance_dag(roles):
    """Validate typed SHA roles and the frozen provenance relationships."""
    if not isinstance(roles, dict):
        raise ReleaseContractError('PROVENANCE_OBJECT_REQUIRED')
    missing = sorted(set(SHA_ROLES) - set(roles))
    if missing:
        raise ReleaseContractError('PROVENANCE_ROLE_MISSING:' + ','.join(missing))
    normalized = {}
    for role in SHA_ROLES:
        value = roles[role]
        if value in (None, 'TBD', 'runtime'):
            normalized[role] = None
        elif role == 'release_critical_tree_attestation':
            if not isinstance(value, str) or not re.fullmatch('[0-9a-f]{64}', value):
                raise ReleaseContractError('RELEASE_CRITICAL_TREE_ATTESTATION_INVALID')
            normalized[role] = value
        else:
            normalized[role] = _sha(value, role)
    equalities = (
        ('reviewed_release_pr_head_sha', 'reviewed_ci_head_sha'),
        ('production_execution_sha', 'release_execution_sha'),
    )
    for left, right in equalities:
        if normalized[left] is not None and normalized[right] is not None:
            if normalized[left] != normalized[right]:
                raise ReleaseContractError('PROVENANCE_EQUALITY_MISMATCH')
    return {
        'schema_version': 1,
        'roles': normalized,
        'nodes': [
            'product_candidate_sha',
            'release_execution_sha',
            'reviewed_release_pr_head_sha',
            'release_merge_sha',
            'production_execution_sha',
        ],
        'equalities': [
            {
                'left': left,
                'right': right,
                'status': 'PASS' if all(normalized[key] for key in (left, right)) else 'TBD',
            }
            for left, right in equalities
        ],
    }


def release_critical_tree_attestation(
    repo,
    release_execution_sha,
    reviewed_release_pr_head_sha,
    paths,
):
    """Prove the release-critical paths are identical at the two frozen heads."""
    _sha(release_execution_sha, 'release_execution_sha')
    _sha(reviewed_release_pr_head_sha, 'reviewed_release_pr_head_sha')
    if not isinstance(paths, (list, tuple)) or not paths:
        raise ReleaseContractError('TREE_ATTESTATION_PATHS_REQUIRED')
    normalized_paths = sorted({str(path) for path in paths if str(path)})
    if not normalized_paths:
        raise ReleaseContractError('TREE_ATTESTATION_PATHS_REQUIRED')

    def tree(sha):
        return subprocess.check_output(
            ['git', 'ls-tree', '-r', '-z', sha, '--', *normalized_paths],
            cwd=repo,
        )

    execution_tree = tree(release_execution_sha)
    reviewed_tree = tree(reviewed_release_pr_head_sha)
    if not execution_tree or not reviewed_tree:
        raise ReleaseContractError('TREE_ATTESTATION_PATH_MISSING')
    if execution_tree != reviewed_tree:
        raise ReleaseContractError('RELEASE_CRITICAL_TREE_MISMATCH')
    return {
        'status': 'PASS',
        'release_execution_sha': release_execution_sha,
        'reviewed_release_pr_head_sha': reviewed_release_pr_head_sha,
        'paths': normalized_paths,
        'tree_sha256': digest(execution_tree),
        'attestation_sha256': digest(execution_tree + reviewed_tree),
    }


def manifest_coverage(planned_targets, resolved_targets, manifest_targets,
                      publish_targets):
    groups = {
        'planned': list(planned_targets),
        'resolved': list(resolved_targets),
        'manifest': list(manifest_targets),
        'publish': list(publish_targets),
    }
    for values in groups.values():
        if any(not isinstance(value, str) or not value for value in values):
            raise ReleaseContractError('MANIFEST_TARGET_INVALID')
        if len(values) != len(set(values)):
            raise ReleaseContractError('MANIFEST_DUPLICATE')
    sets = [set(values) for values in groups.values()]
    if any(value != sets[0] for value in sets[1:]):
        raise ReleaseContractError('MANIFEST_COVERAGE_MISMATCH')
    return {
        'status': 'PASS',
        'planned_targets': len(sets[0]),
        'resolved_targets': len(sets[1]),
        'manifest_targets': len(sets[2]),
        'publish_targets': len(sets[3]),
        'missing': 0,
        'unexpected': 0,
        'duplicate': 0,
        'unresolved': 0,
    }


def classify_remote_error(error):
    status = getattr(error, 'status', None) or getattr(error, 'status_code', None)
    try:
        status = int(status) if status is not None else None
    except (TypeError, ValueError):
        status = None
    text = f'{type(error).__name__}:{error}'.lower()
    if status in RETRYABLE_STATUS or (status is not None and 500 <= status <= 599):
        return RETRYABLE
    if any(term in text for term in NON_RETRYABLE_TERMS):
        return NON_RETRYABLE
    if any(term in text for term in RETRYABLE_TERMS):
        return RETRYABLE
    return NON_RETRYABLE


def bounded_read_retry(call, *, attempts=4, base_delay=0.5, sleeper=time.sleep,
                       reset=None, events=None):
    if attempts < 1:
        raise ValueError('READ_RETRY_ATTEMPTS_INVALID')
    events = events if events is not None else []
    for attempt in range(1, attempts + 1):
        try:
            return call()
        except Exception as exc:
            classification = classify_remote_error(exc)
            events.append({
                'attempt': attempt,
                'classification': classification,
                'error_type': type(exc).__name__,
            })
            if classification != RETRYABLE or attempt == attempts:
                raise
            if reset is not None:
                reset()
            sleeper(base_delay * (2 ** (attempt - 1)))
    raise AssertionError('unreachable')


def reconcile_write(*, journal_state, observed_sha256, before_sha256,
                    after_sha256):
    if observed_sha256 == after_sha256:
        return {'decision': 'ALREADY_COMMITTED', 'state': 'VERIFYING'}
    if observed_sha256 == before_sha256:
        if journal_state not in {
            'PUBLISHING', 'VERIFYING', 'FAILED', 'ARCHIVE_COMPLETE',
        }:
            raise ReleaseContractError('WRITE_RECONCILIATION_STATE_INVALID')
        return {'decision': 'RETRY_ALLOWED', 'state': journal_state}
    raise PartialMutationError('WRITE_RECONCILIATION_UNKNOWN_STATE')


def append_archive_checkpoint(root, *, release_execution_sha, release_id,
                              state, progress, recorded_at):
    _sha(release_execution_sha, 'release_execution_sha')
    if state not in RELEASE_STATES:
        raise ReleaseContractError('ARCHIVE_CHECKPOINT_STATE_INVALID')
    if not isinstance(recorded_at, str) or not recorded_at:
        raise ReleaseContractError('ARCHIVE_CHECKPOINT_TIME_REQUIRED')
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    payload = {
        'schema_version': 1,
        'release_execution_sha': release_execution_sha,
        'release_id': release_id,
        'state': state,
        'progress': progress,
        'recorded_at': recorded_at,
    }
    data = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(',', ':'),
    ).encode() + b'\n'
    with open(root / 'checkpoints.jsonl', 'ab') as stream:
        stream.write(data)
        stream.flush()
        __import__('os').fsync(stream.fileno())
    return payload


def read_archive_checkpoints(root):
    path = Path(root) / 'checkpoints.jsonl'
    if not path.exists():
        return []
    values = []
    for line in path.read_text(encoding='utf-8').splitlines():
        if not line:
            continue
        value = json.loads(line)
        if not isinstance(value, dict) or value.get('schema_version') != 1:
            raise ReleaseContractError('ARCHIVE_CHECKPOINT_INVALID')
        values.append(value)
    return values


def payload(drive,entry):
    if entry.get('mode')=='managed_doc':return drive.get_managed_doc(entry['id'])
    return drive.get(entry['id'])


def write_payload(drive,entry,data):
    if entry.get('mode')=='managed_doc':return drive.put_managed_doc(entry['id'],data)
    return drive.put(entry['id'],data,entry['mime'])


def scan_native_document(drive, entry):
    """Native backups copy the entire Doc, including unmanaged historical text."""
    if entry.get('mode') != 'managed_doc':
        return
    import json
    document = drive.docs.document(entry['id'])
    clean(json.dumps(document, ensure_ascii=False).encode(), 'native/' + entry['name'])


def fingerprint(meta):
    return {k:meta.get(k) for k in ('id','version','modifiedTime','mimeType','parents')}


def snapshot(drive, file_id, mode='binary'):
    before=drive.meta(file_id)
    data=payload(drive,{'id':file_id,'mode':mode})
    after=drive.meta(file_id)
    if fingerprint(before)!=fingerprint(after):
        raise RuntimeError('Remote changed during download')
    return data, after


def prepare(
    drive,
    root,
    release_id,
    entries,
    archive_id,
    status_id,
    dependencies=None,
    carry_forward_artifacts=False,
    closure=None,
    environment='sandbox',
    code_commit=None,
):
    root=Path(root)
    if root.exists(): raise ValueError('Release already exists; resume it instead')
    if not entries: raise ValueError('Empty release')
    ids=[e['id'] for e in entries]
    if len(ids)!=len(set(ids)) or status_id in ids:
        raise ValueError('Duplicate target or status overlaps payload')
    if release_id.startswith('v1.5.4') and closure is None:
        raise ValueError('CANONICAL_CLOSURE_REQUIRED')
    if environment not in {'sandbox', 'production'}:
        raise ValueError('Unknown release environment')
    if environment == 'production' and closure is None:
        import re
        if not isinstance(code_commit, str) or not re.fullmatch('[0-9a-f]{40}', code_commit):
            raise ValueError('PRODUCTION_CODE_COMMIT_REQUIRED')
    plan={'release_id':release_id,'archive_id':archive_id,'status_id':status_id,'entries':[], 'dependencies':dependencies or []}
    plan['environment'] = environment
    if code_commit is not None:
        plan['code_commit'] = code_commit
    if closure is not None:
        plan['closure'] = closure
    clean(__import__('json').dumps(plan).encode(), 'plan.json')
    # Reject secrets before freezing before/candidate copies or remote writes.
    for entry in entries:
        clean(Path(entry['path']).read_bytes(), entry['name'])
    check_dependencies(drive,plan)
    root.mkdir(parents=True)
    try:
        status_data,status_meta=snapshot(drive,status_id)
        clean(status_data, 'status.before.json')
        import json
        status=json.loads(status_data)
        if status.get('state') not in ('COMPLETE','ROLLED_BACK','INITIAL'):
            raise RuntimeError('Another incomplete release needs recovery')
        atomic(root/'status.before.json',status_data)
        plan['status_before_hash']=digest(status_data)
        plan['status_before_meta']=fingerprint(status_meta)
        plan['previous_release_id']=status.get('current_release_id')
        if environment == 'production' and closure is None:
            previous_code_commit = status.get('code_commit')
            import re
            if not isinstance(previous_code_commit, str) or not re.fullmatch('[0-9a-f]{40}', previous_code_commit):
                raise RuntimeError('PREVIOUS_CODE_COMMIT_REQUIRED')
            plan['previous_code_commit'] = previous_code_commit
        if carry_forward_artifacts or closure is not None:
            carried=status.get('artifacts')
            if not isinstance(carried,list) or not carried:
                raise RuntimeError('Previous status has no artifacts to carry forward')
            ids=[item.get('id') for item in carried]
            if any(not item_id for item_id in ids) or len(ids)!=len(set(ids)):
                raise RuntimeError('Previous status artifact list is invalid')
            plan['carry_forward_artifacts']=[
                {'id':item['id'],'name':item['name'],'sha256':item['sha256']}
                for item in carried
            ]
        for i, entry in enumerate(entries):
            name=entry['name']
            if Path(name).name!=name: raise ValueError('Unsafe artifact name')
            new=Path(entry['path']).read_bytes()
            mode=entry.get('mode','binary')
            meta=drive.meta(entry['id'])
            if mode not in ('binary','managed_doc'):raise ValueError('Unknown publication mode')
            allowed_native=mode=='managed_doc' and meta['mimeType']==DOC
            if meta['mimeType']!=entry['mime'] or (meta['mimeType'].startswith('application/vnd.google-apps.') and not allowed_native):
                raise ValueError('MIME mismatch or native object')
            if mode=='managed_doc':
                from .native import BEGIN, END
                text=new.decode()
                if not text.startswith(BEGIN) or not text.endswith(END):raise ValueError('Managed Doc prefix required')
            old,meta=snapshot(drive,entry['id'],mode)
            clean(old, 'before/' + name)
            scan_native_document(drive, dict(entry, mode=mode))
            atomic(root/f'before/{i}',old); atomic(root/f'candidate/{i}',new)
            plan['entries'].append({'id':entry['id'],'name':name,'mime':entry['mime'],'mode':mode,
                'before':f'before/{i}','candidate':f'candidate/{i}',
                'before_hash':digest(old),'after_hash':digest(new),'meta':fingerprint(meta),
                **{k:entry[k] for k in ('staging_parent','publish_parent','logical_key') if k in entry}})
        save(root/'plan.json',plan)
        if closure is not None:
            freeze_closure(drive, root, plan)
            save(root/'plan.json', plan)
            validate_closure(drive, root, plan, candidate=True)
        save(root/'journal.json',{'state':'PREPARED','uploaded':{},'inflight':None})
    except BaseException:
        save(root/'prepare_failed.json',{'state':'PREPARE_FAILED'})
        raise
    return plan


def check_dependencies(drive, plan):
    for dependency in plan.get('dependencies',[]):
        content,meta=snapshot(drive,dependency['id'],dependency.get('mode','binary'))
        clean(content, 'dependency/' + dependency['id'])
        if digest(content)!=dependency['sha256'] or fingerprint(meta)!=dependency['meta']:
            raise RuntimeError('Input dependency changed: '+dependency['id'])


def relocate(drive, entry, destination):
    if not entry.get('staging_parent'):return
    other=entry['staging_parent'] if destination==entry['publish_parent'] else entry['publish_parent']
    parents=drive.meta(entry['id']).get('parents',[])
    if destination in parents:return
    if other not in parents:raise RuntimeError('New object moved outside release ownership')
    drive.move(entry['id'],destination,other)
    if destination not in drive.meta(entry['id']).get('parents',[]):
        raise RuntimeError('Object move verification failed')


def verified_local(root, entry, key):
    data=child(root,entry[key]).read_bytes()
    expected=entry['before_hash' if key=='before' else 'after_hash']
    if digest(data)!=expected: raise RuntimeError('Frozen local artifact was altered')
    return data


def archive_namespace(plan):
    """Production children are isolated by the frozen execution SHA."""
    import json
    import re
    code_commit = (plan.get('closure') or {}).get('code_commit') or plan.get('code_commit')
    if code_commit is not None:
        if not isinstance(code_commit, str) or not re.fullmatch('[0-9a-f]{40}', code_commit):
            raise ValueError('ARCHIVE_CODE_COMMIT_REQUIRED')
        return code_commit
    if plan.get('environment') == 'production':
        raise ValueError('ARCHIVE_CODE_COMMIT_REQUIRED')
    # Legacy sandbox plans have no execution SHA; isolate them by frozen plan bytes.
    return 'plan-' + plan['release_id'] + '-' + digest(json.dumps(plan, sort_keys=True).encode())


def archive_namespace_key(plan):
    """Stable execution identity combines the frozen SHA and release id."""
    return f'{archive_namespace(plan)}:{plan["release_id"]}'


def require_archive_status_unchanged(drive, plan):
    before = fingerprint(drive.meta(plan['status_id']))
    content = drive.get(plan['status_id'])
    after = fingerprint(drive.meta(plan['status_id']))
    if (digest(content) != plan['status_before_hash']
            or before != plan['status_before_meta'] or after != before):
        raise RuntimeError('ARCHIVE_RESUME_PRODUCTION_DRIFT')


def validate_previous_snapshot(plan, previous):
    """Require exactly one verified rollback reference per target/protected object."""
    expected = {e['id']: e['before_hash'] for e in plan['entries']}
    if len(expected) != len(plan['entries']):
        raise RuntimeError('PREVIOUS_SNAPSHOT_INCOMPLETE')
    for item in (plan.get('closure') or {}).get('protected', []):
        expected.setdefault(item['id'], item['sha256'])
    if not isinstance(previous, list) or len(previous) != len(expected):
        raise RuntimeError('PREVIOUS_SNAPSHOT_INCOMPLETE')
    seen = set()
    for item in previous:
        if not isinstance(item, dict):
            raise RuntimeError('PREVIOUS_SNAPSHOT_INCOMPLETE')
        fid = item.get('original_id')
        if (fid not in expected or fid in seen or not item.get('snapshot_id')
                or item.get('sha256') != expected[fid]
                or item.get('hash_scope') not in ('managed_prefix', 'bytes')):
            raise RuntimeError('PREVIOUS_SNAPSHOT_INCOMPLETE')
        seen.add(fid)


def archive_snapshot(drive, root, plan):
    namespace = archive_namespace(plan)
    folder = drive.ensure(plan['archive_id'], plan['release_id'], plan['release_id'], FOLDER)
    index = drive.index_cba_keys(folder)
    previous = []
    for i, e in enumerate(plan['entries']):
        native = e.get('mode') == 'managed_doc'
        native_backup = None
        if native:
            scan_native_document(drive, e)
            native_backup = drive.ensure_copy(
                folder, f'{namespace}:native-before-{i}', e['id'], f'before_{e["name"]}',
                index=index,
            )
            if digest(drive.get_managed_doc(native_backup)) != e['before_hash']:
                raise RuntimeError('Immutable native backup differs')
        for kind in ('before', 'candidate'):
            data = verified_local(root, e, kind)
            fid = drive.ensure(
                folder, f'{namespace}:{kind}-{i}',
                f'{kind}_{e["name"]}' + ('.txt' if native else ''),
                'text/plain' if native else e['mime'], data, index=index,
            )
            if kind == 'before':
                previous.append({
                    'original_id': e['id'], 'snapshot_id': native_backup or fid,
                    'name': e['name'], 'sha256': e['before_hash'],
                    'hash_scope': 'managed_prefix' if native else 'bytes',
                })
    if plan.get('closure'):
        snapshotted = {item['original_id'] for item in previous}
        for i, item in enumerate(plan['closure']['protected']):
            if item['id'] in snapshotted:
                continue
            data = child(root, item['before']).read_bytes()
            if digest(data) != item['sha256']:
                raise RuntimeError('Frozen protected snapshot was altered')
            fid = drive.ensure(folder, f'{namespace}:protected-{i}', f'before_{item["name"]}',
                               'application/octet-stream', data, index=index)
            previous.append({'original_id': item['id'], 'snapshot_id': fid,
                             'name': item['name'], 'sha256': item['sha256'], 'hash_scope': 'bytes'})
            snapshotted.add(item['id'])
        status_before = (root/'status.before.json').read_bytes()
        if digest(status_before) != plan['status_before_hash']:
            raise RuntimeError('Frozen status snapshot was altered')
        drive.ensure(folder, f'{namespace}:status-before', 'status.before.json',
                     'application/json', status_before, index=index)
    drive.ensure(folder, f'{namespace}:plan', 'plan.json', 'application/json',
                 (root/'plan.json').read_bytes(), index=index)
    return previous


def set_status(drive, plan, state, previous_snapshot):
    import json
    if state == 'PUBLISHING':
        validate_previous_snapshot(plan, previous_snapshot)
    artifacts={
        item['id']:dict(item)
        for item in plan.get('carry_forward_artifacts',[])
    }
    artifacts.update({
        e['id']:{'id':e['id'],'name':e['name'],'sha256':e['after_hash']}
        for e in plan['entries']
    })
    status={'state':state,'current_release_id':plan['release_id'] if state=='COMPLETE' else plan['previous_release_id'],
            'pending_release_id':plan['release_id'] if state not in ('COMPLETE','ROLLED_BACK') else None,
            'previous_release_id':plan['previous_release_id'],
            'previous_snapshot':previous_snapshot,
            'artifacts':list(artifacts.values())}
    if plan.get('closure'):
        status['code_commit'] = (
            plan['closure']['code_commit'] if state == 'COMPLETE'
            else plan['closure'].get('previous_code_commit')
        )
    elif plan.get('code_commit'):
        status['code_commit'] = (
            plan['code_commit'] if state == 'COMPLETE'
            else plan.get('previous_code_commit')
        )
        if state == 'ROLLED_BACK':
            status['code_commit'] = plan.get('previous_code_commit')
    content=json.dumps(status,ensure_ascii=False,sort_keys=True).encode()
    clean(content, 'release_status.json')
    drive.put(plan['status_id'],content,'application/json')
    if drive.get(plan['status_id'])!=content: raise RuntimeError('Status readback failed')


def publish(drive, root, single_writer=False):
    root=Path(root)
    if not single_writer: raise RuntimeError('Single-writer maintenance window must be acknowledged')
    with lock(root.parent/'publish.lock'):
        plan,journal=read(root/'plan.json'),read(root/'journal.json')
        security_preflight(root, plan)
        stage('publish',f'release {plan["release_id"]} journal {journal["state"]}')
        if journal['state'] in ('ROLLED_BACK','ROLLING_BACK'):
            raise RuntimeError('Create a new plan after rollback')
        if journal.get('provenance'):
            provenance_dag(plan['provenance'])
        if journal.get('release_critical_tree_attestation'):
            if journal['release_critical_tree_attestation'].get('status') != 'PASS':
                raise PreMutationAbort('TREE_ATTESTATION_REQUIRED')
        if journal.get('manifest_coverage'):
            coverage = journal['manifest_coverage']
            required = {'planned_targets','resolved_targets','manifest_targets','publish_targets'}
            if not required <= set(coverage):
                raise PreMutationAbort('MANIFEST_COVERAGE_REQUIRED')
        if journal['state'] == 'ARCHIVING':
            if journal.get('uploaded') != {} or journal.get('inflight') is not None:
                raise RuntimeError('ARCHIVING_JOURNAL_INVALID')
            require_archive_status_unchanged(drive, plan)
        check_dependencies(drive,plan)
        for e in plan['entries']:
            verified_local(root,e,'before'); verified_local(root,e,'candidate')
        if journal['state']=='COMPLETE':
            current=json.loads(drive.get(plan['status_id']))
            if current.get('state')!='COMPLETE' or current.get('current_release_id')!=plan['release_id']:
                raise RuntimeError('Publication status no longer refers to this complete release')
            verify(drive,root); return journal
        # Preflight every object before any production mutation.
        for e in plan['entries']:
            data,meta=snapshot(drive,e['id'],e.get('mode','binary')); sha=digest(data)
            clean(data, 'readback/' + e['name'])
            scan_native_document(drive, e)
            if e['id'] in journal.get('uploaded',{}) or journal.get('inflight')==e['id']:
                if sha==e['after_hash']: continue
                if sha!=e['before_hash']: raise RuntimeError('Conflict while resuming')
            if fingerprint(meta)!=e['meta'] or sha!=e['before_hash']:
                raise RuntimeError('Remote conflict: '+e['name'])
        if journal['state']=='PREPARED':
            require_archive_status_unchanged(drive, plan)
        if plan.get('closure'):
            validate_closure(drive, root, plan, candidate=True)
        stage('publish','preflight passed')
        state = journal['state']
        pending_states = {'ARCHIVE_COMPLETE','PUBLISHING','VERIFYING'}
        if state not in {'PREPARED','ARCHIVING',*pending_states}:
            status=json.loads(drive.get(plan['status_id']))
            untouched_status=digest(drive.get(plan['status_id']))==plan['status_before_hash']
            if not untouched_status and status.get('pending_release_id')!=plan['release_id'] and status.get('current_release_id')!=plan['release_id']:
                raise RuntimeError('Another publisher changed the release status')
        previous = journal.get('previous_snapshot')
        if state in ('PREPARED', 'ARCHIVING'):
            if state == 'PREPARED':
                validate_state_transition('PREPARED','ARCHIVING')
            journal.update(state='ARCHIVING', uploaded={}, inflight=None)
            save(root/'journal.json', journal)
            previous = archive_snapshot(drive, root, plan)
            validate_previous_snapshot(plan, previous)
            journal['previous_snapshot']=previous
            save(root/'journal.json',journal)
            stage('publish',f'archive snapshot for {len(plan["entries"])} artifacts')
        else:
            if not previous:
                raise PreMutationAbort('ARCHIVE_CHECKPOINT_REQUIRED')
            validate_previous_snapshot(plan, previous)
            stage('publish','verified archive checkpoint reused')
        # Dependency/status drift during a long archive must not change canonical status.
        check_dependencies(drive, plan)
        if state in ('PREPARED', 'ARCHIVING'):
            require_archive_status_unchanged(drive, plan)
            validate_state_transition('ARCHIVING','ARCHIVE_COMPLETE')
            journal['state']='ARCHIVE_COMPLETE'
            save(root/'journal.json',journal)
            version_meta = plan.get('status_before_meta') or {}
            if version_meta.get('modifiedTime'):
                append_archive_checkpoint(
                    root/'archive-checkpoints',
                    release_execution_sha=(
                        (plan.get('closure') or {}).get('code_commit')
                        or plan.get('code_commit') or '0' * 40
                    ),
                    release_id=plan['release_id'],
                    state='ARCHIVE_COMPLETE',
                    progress={'snapshots': len(previous)},
                    recorded_at=version_meta['modifiedTime'],
                )
        try:
            if state not in {'PUBLISHING','VERIFYING'}:
                validate_state_transition('ARCHIVE_COMPLETE','PUBLISHING')
                journal['state']='PUBLISHING'
                save(root/'journal.json',journal)
                set_status(drive,plan,'PUBLISHING',previous)
            else:
                current_status=json.loads(drive.get(plan['status_id']))
                if current_status.get('pending_release_id')!=plan['release_id'] and current_status.get('current_release_id')!=plan['release_id']:
                    raise RuntimeError('Publication status no longer refers to this release')
            for index,e in enumerate(plan['entries'],1):
                stage('publish',f'{index}/{len(plan["entries"])} {e["name"]}')
                data=payload(drive,e)
                if digest(data)==e['after_hash']:
                    journal.setdefault('uploaded',{})[e['id']]=True
                    save(root/'journal.json',journal); continue
                if digest(data)!=e['before_hash']: raise RuntimeError('Remote changed during publication')
                journal['inflight']=e['id']; save(root/'journal.json',journal)
                latest,latest_meta=snapshot(drive,e['id'],e.get('mode','binary'))
                if digest(latest)!=e['before_hash']:raise RuntimeError('Target changed immediately before write')
                try:
                    write_payload(drive,e,verified_local(root,e,'candidate'))
                    clean(payload(drive,e), 'readback/' + e['name'])
                    if digest(payload(drive,e))!=e['after_hash']:
                        raise RuntimeError('Readback mismatch')
                except Exception as exc:
                    if classify_remote_error(exc) != RETRYABLE:
                        raise
                    try:
                        observed=digest(payload(drive,e))
                    except Exception:
                        raise exc
                    decision=reconcile_write(
                        journal_state=journal['state'],
                        observed_sha256=observed,
                        before_sha256=e['before_hash'],
                        after_sha256=e['after_hash'],
                    )
                    if decision['decision'] != 'ALREADY_COMMITTED':
                        raise
                    journal.setdefault('uploaded',{})[e['id']]=True
                    journal['inflight']=None
                    save(root/'journal.json',journal)
                    raise exc
                journal.setdefault('uploaded',{})[e['id']]=True; journal['inflight']=None
                save(root/'journal.json',journal)
            if state == 'PUBLISHING':
                validate_state_transition('PUBLISHING','VERIFYING')
            journal['state']='VERIFYING'; save(root/'journal.json',journal)
            verify(drive,root)
            check_dependencies(drive,plan)
            for e in plan['entries']:relocate(drive,e,e.get('publish_parent'))
            if plan.get('closure'):
                validate_closure(drive, root, plan, candidate=False, relocated=True)
            validate_state_transition('VERIFYING','COMPLETE')
            set_status(drive,plan,'COMPLETE',previous)
            if plan.get('closure'):
                validate_closure(drive, root, plan, candidate=False, relocated=True, final=True)
            journal['state']='COMPLETE'; save(root/'journal.json',journal)
            stage('publish','COMPLETE')
        except BaseException:
            journal['failed_from']=journal.get('state')
            journal['state']='FAILED'; save(root/'journal.json',journal)
            if previous:
                try: set_status(drive,plan,'FAILED',previous)
                except Exception: pass
            raise
        return journal

def security_preflight(root, plan):
    """Scan every frozen byte, including reports and retained rollback evidence."""
    import json
    from .common import child
    clean(json.dumps(plan, sort_keys=True).encode(), 'plan.json')
    labels = {}
    for entry in plan.get('entries', []):
        for field in ('before', 'candidate'):
            if entry.get(field):
                labels[child(root, entry[field]).resolve()] = (
                    entry.get('logical_key') or entry.get('name') or field
                )
    for item in (plan.get('closure') or {}).get('protected', []):
        if item.get('before'):
            labels[child(root, item['before']).resolve()] = item['name']
    for path in sorted(Path(root).rglob('*')):
        if path.is_symlink():
            raise ValueError('RELEASE_SYMLINK_FORBIDDEN')
        if path.is_file():
            clean(
                path.read_bytes(),
                labels.get(path.resolve(), str(path.relative_to(root))),
            )
    if plan['release_id'].startswith('v1.5.4') and not plan.get('closure'):
        raise ValueError('CANONICAL_CLOSURE_REQUIRED')
    if plan.get('environment') == 'production':
        code_commit = (
            plan.get('closure', {}).get('code_commit')
            or plan.get('code_commit')
        )
        if code_commit:
            verify_code_provenance(
                Path(__file__).resolve().parents[2], code_commit,
            )


def verify_code_provenance(repo, code_commit):
    """Require the executing checkout to be the merged commit on the frozen base."""
    import subprocess
    import yaml
    task_path = Path(repo) / 'requirements/REQ-154-CANONSEC-01/task.yaml'
    task = yaml.safe_load(task_path.read_text())
    base = 'refs/remotes/origin/' + task['base_branch']
    def git(*args):
        return subprocess.check_output(['git', *args], cwd=repo, stderr=subprocess.DEVNULL, text=True).strip()
    try:
        if git('rev-parse', 'HEAD') != code_commit or git('status', '--porcelain', '--untracked-files=no'):
            raise ValueError('CODE_MIRROR_NOT_MERGED_COMMIT')
        if subprocess.run(['git', 'merge-base', '--is-ancestor', code_commit, base],
                          cwd=repo, capture_output=True).returncode:
            raise ValueError('CODE_MIRROR_NOT_MERGED_COMMIT')
    except subprocess.CalledProcessError:
        raise ValueError('CODE_MIRROR_NOT_MERGED_COMMIT') from None
    return {'status': 'PASS', 'code_commit': code_commit}


def freeze_closure(drive, root, plan):
    """Freeze real protected bytes, not caller-supplied PASS flags or row counts."""
    import re
    closure = plan['closure']
    required = {'code_commit', 'previous_code_commit', 'baseline_release_id',
                'registry_key', 'manifest_key', 'documents', 'zones', 'protected'}
    if not isinstance(closure, dict) or not required <= set(closure):
        raise ValueError('CLOSURE_CONTRACT_REQUIRED')
    for key in ('code_commit', 'previous_code_commit'):
        if not isinstance(closure[key], str) or not re.fullmatch('[0-9a-f]{40}', closure[key]):
            raise ValueError('CLOSURE_CODE_COMMIT_REQUIRED')
    if closure['baseline_release_id'] != plan['previous_release_id']:
        raise ValueError('BASELINE_REFREEZE_REQUIRED')
    controls = {closure['registry_key'], closure['manifest_key'], *closure['documents'].values()}
    entries = {entry.get('logical_key'): entry for entry in plan['entries']}
    document_keys = set(closure['documents'])
    legacy_documents = {'readme', 'index', 'context_card', 'version'}
    current_documents = legacy_documents | {'current_version_doc'}
    expected_controls = 7 if document_keys == current_documents else 6
    if len(entries) != len(plan['entries']) or None in entries or len(controls) != expected_controls or not controls <= set(entries):
        raise ValueError('CLOSURE_CONTROL_COVERAGE')
    if document_keys not in (legacy_documents, current_documents):
        raise ValueError('CLOSURE_DOCUMENT_COVERAGE')
    protected = closure['protected']
    if not isinstance(protected, list) or not protected:
        raise ValueError('PROTECTED_BASELINE_REQUIRED')
    kinds = {item.get('kind') for item in protected}
    if not {'master', 'six_table', 'source_registry', 'evidence'} <= kinds:
        raise ValueError('PROTECTED_BASELINE_COVERAGE')
    if len({item.get('id') for item in protected}) != len(protected):
        raise ValueError('PROTECTED_BASELINE_DUPLICATE')
    for index, item in enumerate(protected):
        if (not item.get('id') or not item.get('name')
                or not re.fullmatch('[0-9a-f]{64}', item.get('sha256', ''))):
            raise ValueError('PROTECTED_BASELINE_INVALID')
        data, meta = snapshot(drive, item['id'], item.get('mode', 'binary'))
        clean(data, item['name'])
        if digest(data) != item['sha256']:
            raise ValueError('BASELINE_REFREEZE_REQUIRED')
        item['before'] = f'protected/{index}'
        item['meta'] = fingerprint(meta)
        atomic(root / item['before'], data)
    security_preflight(root, plan)


def verify_protected(drive, root, plan):
    """All business/source/evidence bytes remain identical throughout publication."""
    protected = plan['closure']['protected']
    candidate = {entry['id']: entry for entry in plan['entries']}
    for item in protected:
        before = child(root, item['before']).read_bytes()
        clean(before, item['name'])
        now = payload(drive, item)
        clean(now, item['name'])
        if digest(before) != item['sha256'] or digest(now) != item['sha256']:
            raise ValueError('ZERO_BUSINESS_FACT_DELTA')
        if item['id'] in candidate and candidate[item['id']]['after_hash'] != item['sha256']:
            raise ValueError('ZERO_BUSINESS_FACT_DELTA')
    return {'status': 'PASS', 'protected_artifacts': len(protected)}


def _rows(data, product):
    import csv
    import io
    import json
    key = product['artifact_key']
    if key.endswith('.jsonl'):
        return [json.loads(line) for line in data.decode('utf-8-sig').splitlines() if line.strip()]
    if key.endswith('.csv'):
        return list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
    if key.endswith('.xlsx'):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            selector = product.get('selector')
            if selector is None and len(book.sheetnames) == 1:
                selector = book.sheetnames[0]
            if selector not in book.sheetnames:
                raise ValueError('PRODUCT_WORKSHEET_UNRESOLVED')
            rows = iter(book[selector].values)
            columns = next(rows)
            if any(not isinstance(value, str) or not value for value in columns) or len(columns) != len(set(columns)):
                raise ValueError('PRODUCT_HEADER_INVALID')
            return [dict(zip(columns, row)) for row in rows if any(value is not None for value in row)]
        finally:
            book.close()
    raise ValueError('PRODUCT_FORMAT_UNSUPPORTED')


def _inventory(drive, zones):
    """Read only configured current/evidence trees; never crawl the archive."""
    from .current_state import NON_CURRENT
    items, seen, parents = {}, set(), {}
    stopped = set(zones['history']) | set(zones['staging'])
    queue = list(zones['current']) + list(zones['evidence'])
    while queue:
        folder = queue.pop()
        if folder in seen or folder in stopped:
            continue
        seen.add(folder)
        for item in drive.list(folder):
            items[item['id']] = item
            if item['mimeType'] == FOLDER:
                parents[item['id']] = item.get('parents', [])
                # Staging/history identities are boundaries even when nested under root.
                if item['id'] not in stopped and item.get('artifact_role') not in NON_CURRENT:
                    queue.append(item['id'])
    return list(items.values()), dict(zones, folders=parents)


def validate_closure(drive, root, plan, *, candidate, relocated=False, final=False):
    import copy
    import csv
    import io
    import json
    from .canonical_registry import (
        load_registry, manifest_index, reconcile_compatibility, validate_registry,
    )
    from .current_state import audit_current_history, validate_current_state
    closure = plan['closure']
    entries = {entry['logical_key']: entry for entry in plan['entries']}
    verify_protected(drive, root, plan)

    def content(key):
        entry = entries[key]
        data = verified_local(root, entry, 'candidate') if candidate else payload(drive, entry)
        clean(data, entry['name'])
        return data

    registry = load_registry(content(closure['registry_key']))
    manifest = content(closure['manifest_key'])
    by_key = manifest_index(manifest)
    protected = {item['id']: item for item in closure['protected']}
    sources = [item for item in protected.values() if item['kind'] == 'source_registry']
    if len(sources) != 1:
        raise ValueError('SOURCE_REGISTRY_BASELINE_REQUIRED')
    source_rows = list(csv.DictReader(io.StringIO(
        child(root, sources[0]['before']).read_bytes().decode('utf-8-sig')
    )))
    products = validate_registry(registry, manifest, source_rows)
    status = {
        'state': 'COMPLETE', 'current_release_id': plan['release_id'],
        'code_commit': closure['code_commit'],
    }
    if final:
        status = json.loads(drive.get(plan['status_id']))
    documents = {role: content(key).decode('utf-8') for role, key in closure['documents'].items()}
    validate_current_state(status, registry, manifest, documents, closure.get('counts'), closure.get('blockers'))
    # Transport hashes stay in manifest. Resolve and verify all current products.
    records = {}
    for pid, product in products.items():
        if not product['current_eligible']:
            continue
        key = product['artifact_key']
        row = by_key[key]
        fid = row.get('drive_file_id') or row.get('id')
        if product['authority'] == 'canonical' and fid not in protected:
            raise ValueError('CANONICAL_FACT_BASELINE_REQUIRED')
        if key in entries:
            if entries[key]['id'] != fid:
                raise ValueError('MANIFEST_TARGET_MISMATCH')
            data = content(key)
        else:
            if fid not in protected:
                raise ValueError('PRODUCT_BASELINE_REQUIRED')
            data = payload(drive, protected[fid])
        clean(data, key)
        if digest(data) != (row.get('content_hash') or row.get('sha256')):
            raise ValueError('MANIFEST_PRODUCT_HASH_MISMATCH')
        records[pid] = _rows(data, product)
    for pid, product in products.items():
        if product['current_eligible'] and product['authority'] in {'compatibility', 'derived'}:
            annotation = (product.get('projection') or {}).get('legacy_annotations')
            if annotation:
                # Supplementary legacy context is parsed from pinned original
                # evidence, never from a compatibility output or another date.
                key = annotation['artifact_key']
                row = by_key.get(key)
                fid = (row.get('drive_file_id') or row.get('id')) if row else None
                if fid not in protected or protected[fid]['kind'] != 'evidence':
                    raise ValueError('LEGACY_ANNOTATION_EVIDENCE_NOT_PROTECTED')
                source = protected[fid]
                if (row.get('content_hash') or row.get('sha256')) != source['sha256']:
                    raise ValueError('LEGACY_ANNOTATION_HASH_MISMATCH')
                from .adapters import midseason_md
                from .aliases import Clubs
                if annotation['parser'] != 'midseason_md':
                    raise ValueError('LEGACY_ANNOTATION_PARSER_FORBIDDEN')
                result = midseason_md.extract(
                    child(root, source['before']),
                    product['projection']['scope']['domestic']['season'],
                    {'id': fid, 'url': f'https://drive.google.com/file/d/{fid}/view', 'type': 'gdrive'},
                    clubs=Clubs(Path(__file__).resolve().parents[2]), strict=True,
                )
                records[key] = result['events']
            reconcile_compatibility(registry, pid, records, records[pid])
    # Verify every changed transport record, including the new controls.
    for key, entry in entries.items():
        row = by_key.get(key)
        if row is None or (row.get('drive_file_id') or row.get('id')) != entry['id']:
            raise ValueError('MANIFEST_TARGET_MISMATCH')
        if key != closure['manifest_key'] and (row.get('content_hash') or row.get('sha256')) != entry['after_hash']:
            raise ValueError('MANIFEST_TARGET_HASH_MISMATCH')
    items, zones = _inventory(drive, closure['zones'])
    # Before relocation, validate the intended current location as well as bytes.
    by_id = {item['id']: copy.deepcopy(item) for item in items}
    zone_roots = [
        *zones['current'], *zones['history'], *zones['staging'], *zones['evidence'],
    ]
    for entry in plan['entries']:
        item = copy.deepcopy(drive.meta(entry['id']))
        target = entry.get('publish_parent')
        if target and _under([target], zone_roots, zones.get('folders', {})):
            if not relocated:
                item['parents'] = [target]
        elif not _under(item.get('parents', []), zone_roots, zones.get('folders', {})):
            continue
        by_id[entry['id']] = item
    audit_current_history(list(by_id.values()), zones, manifest, registry)
    # Evidence coverage must be complete for the configured evidence trees.
    evidence_ids = {
        item['id'] for item in items
        if item['mimeType'] != FOLDER and _under(item.get('parents', []), zones['evidence'], zones.get('folders', {}))
    }
    evidence_ids.update(
        entry['id'] for entry in plan['entries']
        if entry.get('publish_parent')
        and _under([entry['publish_parent']], zones['evidence'], zones.get('folders', {}))
    )
    if evidence_ids != {item['id'] for item in protected.values() if item['kind'] == 'evidence'}:
        raise ValueError('EVIDENCE_BASELINE_COVERAGE')
    return {'status': 'PASS', 'release_id': plan['release_id']}


def _under(parents, roots, folders):
    seen, queue = set(), list(parents)
    while queue:
        item = queue.pop()
        if item in roots:
            return True
        if item not in seen:
            seen.add(item)
            queue.extend(folders.get(item, []))
    return False


def verify(drive, root):
    plan=read(Path(root)/'plan.json')
    security_preflight(Path(root), plan)
    for e in plan['entries']:
        content = payload(drive,e)
        clean(content, 'readback/' + e['name'])
        scan_native_document(drive, e)
        if drive.meta(e['id'])['mimeType']!=e['mime'] or digest(content)!=e['after_hash']:
            raise RuntimeError('Remote verification failed: '+e['name'])
    if plan.get('closure'):
        validate_closure(drive, Path(root), plan, candidate=False)
    return {'verified':len(plan['entries'])}


def restore(drive, root, single_writer=False):
    root=Path(root)
    if not single_writer: raise RuntimeError('Single-writer maintenance window required')
    with lock(root.parent/'publish.lock'):
        plan,journal=read(root/'plan.json'),read(root/'journal.json')
        security_preflight(root, plan)
        stage('restore',f'release {plan["release_id"]} from {journal["state"]}')
        if digest((root/'status.before.json').read_bytes())!=plan['status_before_hash']:
            raise RuntimeError('Frozen previous status was altered')
        import json
        status=json.loads(drive.get(plan['status_id']))
        already=(status.get('state')=='ROLLED_BACK' and status.get('rolled_back_release_id')==plan['release_id'])
        if not already and status.get('pending_release_id')!=plan['release_id'] and status.get('current_release_id')!=plan['release_id']:
            raise RuntimeError('Cannot roll back another release')
        for e in plan['entries']:
            verified_local(root,e,'before')
            clean(payload(drive,e), 'restore/' + e['name'])
            if digest(payload(drive,e)) not in (e['before_hash'],e['after_hash']):
                raise RuntimeError('Rollback would overwrite an external edit')
        if already:
            for e in plan['entries']:
                if digest(payload(drive,e))!=e['before_hash']:raise RuntimeError('Restored object changed externally')
                if e.get('staging_parent') and e['staging_parent'] not in drive.meta(e['id']).get('parents',[]):
                    raise RuntimeError('Restored object parent changed')
            journal['state']='ROLLED_BACK';save(root/'journal.json',journal);return journal
        journal['state']='ROLLING_BACK';save(root/'journal.json',journal)
        set_status(drive,plan,'ROLLING_BACK',journal.get('previous_snapshot',[]))
        stage('restore','ROLLING_BACK published; restoring original bytes')
        for index,e in enumerate(reversed(plan['entries']),1):
            stage('restore',f'{index}/{len(plan["entries"])} {e["name"]}')
            old=verified_local(root,e,'before')
            now=payload(drive,e)
            if digest(now) not in (e['before_hash'],e['after_hash']):raise RuntimeError('Target changed during rollback')
            if digest(now)!=e['before_hash']: write_payload(drive,e,old)
            if digest(payload(drive,e))!=e['before_hash']: raise RuntimeError('Rollback readback failed')
            clean(payload(drive,e), 'restore/' + e['name'])
            relocate(drive,e,e.get('staging_parent'))
        previous_status=read(root/'status.before.json')
        if plan.get('closure'):
            verify_protected(drive, root, plan)
        previous_status['state']='ROLLED_BACK'
        previous_status['rolled_back_release_id']=plan['release_id']
        data=json.dumps(previous_status,ensure_ascii=False,sort_keys=True).encode()
        drive.put(plan['status_id'],data,'application/json')
        if drive.get(plan['status_id'])!=data: raise RuntimeError('Rollback status verification failed')
        journal['state']='ROLLED_BACK'; save(root/'journal.json',journal)
        stage('restore','ROLLED_BACK')
        return journal


PERFORMANCE_METRIC_FIELDS = (
    "normal_codex_calls",
    "CI_repair_codex_calls",
    "local_exception_codex_calls",
    "CI_wall_clock_seconds",
    "local_prepare_wall_clock_seconds",
    "publish_wall_clock_seconds",
    "context_expansion_count",
    "refreeze_count",
    "development_sync_count",
    "web_merge_review_cycles",
    "WEB05_not_ready_attempts",
    "contract_semantic_defect_count",
    "production_resilience_defect_count",
    "recovery_cycle_count",
    "pre_mutation_abort_count",
    "archive_reuse_count",
    "transport_reconnect_count",
    "partial_write_failure_count",
    "rollback_count",
    "consumer_baseline_failure_layer_counts",
    "watcher_qualified_cycles",
)


def validate_performance_metrics(metrics):
    if not isinstance(metrics, dict):
        raise ReleaseContractError('PERFORMANCE_METRICS_OBJECT_REQUIRED')
    missing = sorted(set(PERFORMANCE_METRIC_FIELDS) - set(metrics))
    if missing:
        raise ReleaseContractError('PERFORMANCE_METRICS_MISSING:' + ','.join(missing))
    for key in PERFORMANCE_METRIC_FIELDS:
        value = metrics[key]
        if value == 'NOT_AVAILABLE':
            continue
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ReleaseContractError('PERFORMANCE_METRIC_VALUE_INVALID')
        if value < 0:
            raise ReleaseContractError('PERFORMANCE_METRIC_VALUE_INVALID')
    return {'status': 'PASS', 'metrics': len(PERFORMANCE_METRIC_FIELDS)}
