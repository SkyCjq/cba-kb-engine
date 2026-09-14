"""Read-only v1.5.3 event reconciliation and cross-table QA.

The command expects frozen local workbook copies.  It does not call Drive and
does not publish production facts.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from openpyxl import load_workbook
from cba_kb.aliases import Clubs
from cba_kb.event_closure import (
    cross_table_qa,
    reconcile_legacy_events,
    semantic_fingerprint,
    upgrade_records,
    validate_event_closure,
    write_reconciliation_reports,
)
from cba_kb.registration_domain import TABLES, parse_domestic_movement


DOMAIN_SHEETS = {
    'domestic_registrations': 'domestic_registrations',
    'domestic_transaction_windows': 'domestic_transaction_windows',
    'registration_status_events': 'registration_status_events',
    'foreign_registration_snapshots': 'foreign_registration_snapshots',
    'foreign_priority_right_snapshots': 'foreign_right_snapshots',
    'foreign_priority_right_transactions': 'foreign_right_transactions',
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def sheet_rows(path, sheet):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise ValueError(f'Missing sheet {sheet!r} in {path}')
        values = book[sheet].iter_rows(values_only=True)
        try:
            headers = list(next(values))
        except StopIteration:
            raise ValueError(f'Empty sheet {sheet!r} in {path}') from None
        rows = [
            {header: row[index] if index < len(row) else None
             for index, header in enumerate(headers)}
            for row in values
        ]
        return headers, rows
    finally:
        book.close()


def load_domain(path):
    records = {}
    headers = {}
    for table, sheet in DOMAIN_SHEETS.items():
        table_headers, rows = sheet_rows(path, sheet)
        records[table] = rows
        headers[table] = table_headers
    return upgrade_records(records), headers


def freeze_report(domain_path, legacy_path, master_path, records, legacy_rows, master_rows,
                  source_windows=None):
    events = records.get('registration_status_events') or []
    windows = records.get('domestic_transaction_windows') or []
    return {
        'domain_workbook': str(domain_path),
        'domain_workbook_sha256': sha256(domain_path),
        'legacy_events_workbook': str(legacy_path),
        'legacy_events_sha256': sha256(legacy_path),
        'master_workbook': str(master_path) if master_path else None,
        'master_sha256': sha256(master_path) if master_path else None,
        'table_counts': {name: len(rows) for name, rows in records.items()},
        'legacy_row_count': len(legacy_rows),
        'master_row_count': len(master_rows),
        'event_domain_counts': dict(sorted(Counter(
            row.get('event_domain') for row in events).items())),
        'event_status_counts': dict(sorted(Counter(
            row.get('event_status') for row in events).items())),
        'windows_by_season': dict(sorted(Counter(
            row.get('season') for row in windows).items())),
        'source_window_closure': source_windows,
        'semantic_sha256': semantic_fingerprint(records),
    }


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + '\n',
                    encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--domain', type=Path, required=True)
    parser.add_argument('--legacy', type=Path, required=True)
    parser.add_argument('--master', type=Path)
    parser.add_argument('--domestic-source', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--expected-legacy-rows', type=int)
    parser.add_argument('--expected-edges', type=int)
    parser.add_argument('--expected-event-domain', action='append', default=[])
    parser.add_argument('--require-window-closure', action='store_true')
    parser.add_argument('--allow-unmapped', action='store_true')
    args = parser.parse_args()

    records, headers = load_domain(args.domain)
    _, legacy_rows = sheet_rows(args.legacy, 'EVENTS')
    master_rows = sheet_rows(args.master, 'MASTER')[1] if args.master else []
    validate_event_closure(records)
    source_windows = None
    if args.domestic_source:
        source = parse_domestic_movement(
            args.domestic_source.read_text(),
            {
                'id': 'v1.5.3-domestic-source',
                'url': f'file://{args.domestic_source.resolve()}',
                'type': 'local_file',
            },
            Clubs(Path(__file__).resolve().parents[1]),
        )
        source_windows = [
            {
                'window_key': row['window_key'],
                'season': row['season'],
                'window_no': row['window_no'],
                'window_start': row['window_start'],
                'window_end': row['window_end'],
                'completeness_status': row['completeness_status'],
                'source_page_or_row': row['source_page_or_row'],
            }
            for row in source['domestic_transaction_windows']
        ]

    edges, summary = reconcile_legacy_events(
        legacy_rows, records, legacy_source_id=sha256(args.legacy)[:16],
    )
    qa = cross_table_qa(records, master_rows, Clubs(Path(__file__).resolve().parents[1]))
    freeze = freeze_report(
        args.domain, args.legacy, args.master, records, legacy_rows, master_rows,
        source_windows,
    )

    if args.expected_legacy_rows is not None and summary['legacy_row_count'] != args.expected_legacy_rows:
        raise ValueError('Legacy row control changed')
    if args.expected_edges is not None and summary['mapping_edge_count'] != args.expected_edges:
        raise ValueError('Mapping-edge control changed')
    for domain in args.expected_event_domain:
        if not freeze['event_domain_counts'].get(domain):
            raise ValueError(f'Expected event domain is missing: {domain}')
    if args.require_window_closure:
        counts = Counter(row['season'] for row in (source_windows or []))
        if counts.get('2024-2025') != 3 or counts.get('2025-2026') != 3:
            raise ValueError('Domestic window closure requires 3 windows in each focus season')
    if not args.allow_unmapped and summary['unmapped_legacy_count']:
        raise ValueError('Legacy reconciliation has unmapped rows')
    if summary['unexplained_conflict_count']:
        raise ValueError('Legacy reconciliation has unexplained conflicts')

    csv_path, reconciliation_summary = write_reconciliation_reports(edges, summary, args.output)
    freeze_path = args.output / 'v1.5.3-phase0-freeze.json'
    qa_path = args.output / 'v1.5.3-cross-table-qa.json'
    schema_path = args.output / 'v1.5.3-event-schema.json'
    write_json(freeze_path, freeze)
    write_json(qa_path, qa)
    write_json(schema_path, {
        'headers': headers,
        'upgraded_event_headers': list(TABLES['registration_status_events']),
    })
    print(json.dumps({
        'freeze': freeze,
        'reconciliation': summary,
        'qa': qa,
        'outputs': [str(csv_path), str(reconciliation_summary), str(freeze_path),
                    str(qa_path), str(schema_path)],
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
