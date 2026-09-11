"""Read-only source acceptance for the implemented DRAFT-2 table scope.

Inputs stay outside Git. This does not accept the unfinished DRAFT-2 scope or
publish production. Counts are controls for the pinned merged source revisions.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import sys
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from openpyxl import load_workbook
from cba_kb.aliases import Clubs
from cba_kb.adapters import foreign_xlsx
from cba_kb.domain_xlsx import SHEET_NAMES, write
from cba_kb.registration_domain import TABLES, parse_domestic_movement, parse_foreign_rights, validate_domain

ROOT = Path(__file__).resolve().parents[1]
SOURCE_IDS = {'domestic': '1dOvVJBVahuyq0L2fI4WRhRy7QJOxdOUR',
              'foreign': '1eGrMDep9YMfPNy66M6-6Jwm9iYaF_gQj'}
COUNTS = dict(zip(TABLES, (46, 4, 72, 244, 227, 7)))
SNAPSHOTS = {'2020-2021': (45, '2021-04-11'), '2022-2023': (52, '2023-04-05'),
             '2024-2025': (73, '2025-03-31'), '2025-2026': (74, '2026-04-24')}
ANNUAL_RIGHTS = dict(zip(('2019-2020','2020-2021','2021-2022','2022-2023',
                         '2023-2024','2024-2025','2025-2026','2026-2027'),
                        (12,23,27,17,15,18,27,27)))


def check(condition, detail):
    if not condition:
        raise ValueError(detail)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--domestic-source', type=Path, required=True)
    parser.add_argument('--rights-source', type=Path, required=True)
    parser.add_argument('--foreign-xlsx', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    check(not args.output.exists(), 'Choose a new verification output directory')
    paths = {'domestic': args.domestic_source, 'foreign': args.rights_source}
    texts = {name: path.read_text() for name, path in paths.items()}
    sources = {name: {'id': id_, 'url': f'https://drive.google.com/file/d/{id_}/view'}
               for name, id_ in SOURCE_IDS.items()}
    clubs = Clubs(ROOT)
    domestic = parse_domestic_movement(texts['domestic'], sources['domestic'], clubs)
    foreign = parse_foreign_rights(texts['foreign'], sources['foreign'], clubs)
    records = {name: domestic.get(name, []) + foreign.get(name, []) for name in TABLES}
    summary = validate_domain(records)
    check({name: item['rows'] for name, item in summary.items()} == COUNTS, 'Six-table count mismatch')
    check(not clubs.report()['unresolved'], 'Unresolved club aliases')
    check(Counter(row['season'] for row in records['foreign_registration_snapshots']) ==
          {s: value[0] for s, value in SNAPSHOTS.items()}, 'Registration season count mismatch')
    for row in records['foreign_registration_snapshots']:
        check(row['snapshot_as_of'] == SNAPSHOTS[row['season']][1], 'Registration snapshot date mismatch')
        check(not re.search(r'[\u3400-\u9fff/]', row['name_en_raw']), 'Mixed registration names')
    check({item['season']: item['rows'] for item in foreign['report']['tables']
           if item['kind'] == 'renewal_status'} == ANNUAL_RIGHTS, 'Annual rights season/count mismatch')
    check(sum(item['rows'] for item in foreign['report']['tables']
              if item['kind'] == 'rights_transaction_status') == 61, 'Dynamic rights count mismatch')
    by_player = {row['player_name_zh']: row for row in domestic['registration_status_events']}
    check(by_player['黄荣奇']['club_announcement_date'] == '2025-11-24', 'Announcement date lost')
    check(by_player['黄荣奇']['event_date'] is None, 'Announcement promoted to event date')
    check(by_player['罗汉琛']['event_date'] == '2026-01-16', 'Window label hid event date')
    check(by_player['罗汉琛']['registration_method'] == '自由球员认领', 'Research label used as official method')
    source_lines = {SOURCE_IDS[key]: text.splitlines() for key, text in texts.items()}
    for table, rows in records.items():
        for row in rows:
            raw_field = next(key for key in TABLES[table] if key.startswith('raw_'))
            raw_lines = row[raw_field].splitlines()
            original = source_lines[row['source_file_id']]
            locators = [int(n) for n in re.findall(r'Markdown line (\d+)', row['source_page_or_row'])]
            check(locators and all(original[n-1] in raw_lines for n in locators), 'Incorrect source locator')
            check(all(line in original for line in raw_lines), 'Raw evidence reconstructed or lost')
            check(row['source_authority'] in {'B2_secondary_cross_check','C_unverified'}, 'Source authority promoted')
    original = foreign_xlsx.extract(args.foreign_xlsx, '2024-2025',
                                   {'id': 'local:foreign_original.xlsx', 'type': 'local'}, clubs=Clubs(ROOT))
    legacy_counts = {key: original['report'][key] for key in
                     ('teams','snapshots','events','observed_registrations')}
    check(legacy_counts == {'teams':20,'snapshots':73,'events':59,'observed_registrations':132},
          'Legacy foreign XLSX compatibility mismatch')
    with TemporaryDirectory() as temporary:
        first, second = (Path(temporary) / name for name in ('first.xlsx','second.xlsx'))
        write(first, records)
        write(second, records)
        check(sha(first) == sha(second), 'Nondeterministic domain workbook')
        book = load_workbook(first, read_only=True, data_only=False)
        try:
            check(book.sheetnames == list(SHEET_NAMES.values()), 'Sheet topology mismatch')
            for table, headers in TABLES.items():
                values = list(book[SHEET_NAMES[table]].values)
                expected = [tuple(headers)]
                for record in records[table]:
                    expected.append(tuple(json.dumps(record[h], ensure_ascii=False, sort_keys=True)
                                          if isinstance(record[h], (list,tuple,dict)) else record[h] or None
                                          if isinstance(record[h], str) else record[h] for h in headers))
                check(values == expected, f'Workbook round-trip mismatch: {table}')
        finally:
            book.close()
        args.output.mkdir(parents=True)
        workbook = args.output / 'domain.xlsx'
        workbook.write_bytes(first.read_bytes())
    report = {'status':'PASS_IMPLEMENTED_TABLE_SCOPE', 'production_eligible':False,
              'version':'1.5.2.dev0', 'counts':COUNTS, 'summary':summary,
              'registration_snapshots_by_season':SNAPSHOTS, 'annual_rights_by_target_season':ANNUAL_RIGHTS,
              'source_hashes':{name:sha(path) for name,path in paths.items()},
              'legacy_foreign_xlsx_sha256':sha(args.foreign_xlsx), 'legacy_foreign_counts':legacy_counts,
              'workbook_sha256':sha(workbook), 'unresolved_clubs':[],
              'checks':['strict aliases','source row locators and literal evidence','season counts and dates',
                        'name separation','announcement/date precision','legacy XLSX compatibility',
                        'six-tab complete cell round-trip','deterministic XLSX'],
              'scope_note':'Table-scope verification only. See docs/CBA-KB_v1.5.2.md for all unfinished requirements.'}
    (args.output / 'source-acceptance.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
