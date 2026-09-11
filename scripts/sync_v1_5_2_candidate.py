"""Create or verify the pinned DRAFT-2 candidate in Google Drive.

The command validates the local workbook against its acceptance report before
any mutation. It only creates a new native Sheet under the requested folder;
an existing object is reused only when its six-tab values match. This is a
candidate sync, not a production release.
"""
import argparse
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from openpyxl import load_workbook
from cba_kb.common import atomic
from cba_kb.drive import Drive
from cba_kb.domain_xlsx import SHEET_NAMES
from cba_kb.registration_domain import TABLES


XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
SHEET = 'application/vnd.google-apps.spreadsheet'
SHEET_ORDER = list(SHEET_NAMES.values())
LEGACY_EVENT_HEADERS = (
    'event_key','event_domain','event_type','event_status','transaction_group_id',
    'season','player_type','club_id','club_source_name','player_name_zh',
    'player_name_en_raw','player_name_en_normalized','from_club_id',
    'from_club_source_name','to_club_id','to_club_source_name','event_date',
    'event_date_status','date_year_inferred','registration_submission_date',
    'registration_submission_date_status','registration_completed_date',
    'registration_completed_date_status','club_announcement_date',
    'registration_window_deadline','registration_method','contract_category',
    'contract_term_official','notes','supersedes_event_key','correction_reason',
    'source_file_id','source_url_primary','source_url_secondary','source_page_or_row',
    'source_type','extraction_method','verification_status','source_authority',
    'verification_raw','source_authority_raw','raw_event_text',
)


def digest(data):
    return hashlib.sha256(data).hexdigest()


def normalize(value):
    if value is None or value == '':
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def pad_row(row, width):
    values = [normalize(value) for value in row]
    return values + [None] * (width - len(values))


def trim_empty_rows(rows):
    end = 0
    for index, row in enumerate(rows, 1):
        if any(normalize(value) is not None for value in row):
            end = index
    return rows[:end]


def book_values(book):
    if book.sheetnames != SHEET_ORDER:
        raise ValueError(f'Candidate sheet order mismatch: {book.sheetnames}')
    values = {}
    for sheet_name in book.sheetnames:
        rows = [list(row) for row in book[sheet_name].iter_rows(values_only=True)]
        if not rows:
            raise ValueError(f'Candidate sheet is empty: {sheet_name}')
        width = len(rows[0])
        if any(len(row) > width for row in rows):
            raise ValueError(f'Candidate row exceeds header width: {sheet_name}')
        values[sheet_name] = [row + [None] * (width - len(row)) for row in rows]
    return values


def workbook_values(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        return book_values(book)
    finally:
        book.close()


def local_checks(candidate, acceptance):
    report = json.loads(Path(acceptance).read_text())
    workbook_sha = digest(Path(candidate).read_bytes())
    checks = {
        'workbook_sha256': workbook_sha,
        'expected_workbook_sha256': report.get('workbook_sha256'),
        'production_eligible': report.get('production_eligible'),
        'version': report.get('version'),
        'status': report.get('status'),
    }
    if report.get('workbook_sha256') != workbook_sha:
        raise ValueError('Candidate workbook does not match the acceptance report')
    if report.get('production_eligible') is not False:
        raise ValueError('Candidate sync requires production_eligible=false')
    expected_counts = report.get('counts')
    if not isinstance(expected_counts, dict):
        raise ValueError('Acceptance report has no table counts')

    values = workbook_values(candidate)
    for table, sheet_name in SHEET_NAMES.items():
        rows = values[sheet_name]
        headers = list(TABLES[table])
        if (report.get('version', '').startswith('1.5.2')
                and table == 'registration_status_events'):
            headers = list(LEGACY_EVENT_HEADERS)
        if rows[0] != headers:
            raise ValueError(f'Candidate headers mismatch: {sheet_name}')
        if len(rows) - 1 != expected_counts.get(table):
            raise ValueError(f'Candidate count mismatch: {sheet_name}')
    checks['counts'] = {table: len(values[sheet]) - 1 for table, sheet in SHEET_NAMES.items()}
    checks['rows_total'] = sum(checks['counts'].values())
    return values, checks


def sheet_values(drive, spreadsheet_id):
    from googleapiclient.http import MediaIoBaseDownload
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(
        buffer, drive.api.files().export_media(fileId=spreadsheet_id, mimeType=XLSX))
    done = False
    while not done:
        _, done = downloader.next_chunk(num_retries=0)
    exported = buffer.getvalue()
    book = load_workbook(io.BytesIO(exported), read_only=True, data_only=False)
    try:
        return book_values(book), exported
    finally:
        book.close()


def compare_values(local, remote):
    for sheet_name in SHEET_ORDER:
        local_rows = trim_empty_rows(local[sheet_name])
        remote_rows = trim_empty_rows(remote[sheet_name])
        if len(local_rows) != len(remote_rows):
            raise ValueError(f'Drive row count mismatch: {sheet_name}')
        width = len(local_rows[0])
        for index, (local_row, remote_row) in enumerate(zip(local_rows, remote_rows), 1):
            expected = pad_row(local_row, width)
            actual = pad_row(remote_row, width)
            if expected != actual:
                raise ValueError(f'Drive value mismatch: {sheet_name} row {index}')


def semantic_sha(values):
    payload = json.dumps(values, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return digest(payload.encode())


def ensure_sheet(drive, folder_id, key, name, candidate):
    matches = [item for item in drive.list(folder_id)
               if item.get('appProperties', {}).get('cba_key') == key]
    if len(matches) > 1:
        raise RuntimeError('Duplicate candidate Sheet key')
    if matches:
        if matches[0]['mimeType'] != SHEET:
            raise ValueError('Existing candidate key is not a native Google Sheet')
        return drive.meta(matches[0]['id']), False

    from googleapiclient.http import MediaIoBaseUpload
    media = MediaIoBaseUpload(io.BytesIO(Path(candidate).read_bytes()), mimetype=XLSX)
    body = {
        'name': name,
        'mimeType': SHEET,
        'parents': [folder_id],
        'appProperties': {'cba_key': key},
    }
    created = drive.api.files().create(
        body=body,
        media_body=media,
        fields='id',
        supportsAllDrives=True,
    ).execute(num_retries=0)
    return drive.meta(created['id']), True


def ensure_attachment(drive, folder_id, key, name, mime, path):
    content = Path(path).read_bytes()
    matches = [item for item in drive.list(folder_id)
               if item.get('appProperties', {}).get('cba_key') == key]
    if len(matches) > 1:
        raise RuntimeError(f'Duplicate attachment key: {key}')
    if matches:
        item = matches[0]
        if item['mimeType'] != mime:
            raise ValueError(f'Attachment MIME mismatch: {key}')
        if drive.get(item['id']) != content:
            raise RuntimeError(f'Immutable attachment differs: {key}')
        return item, False

    from googleapiclient.http import MediaIoBaseUpload
    body = {
        'name': name,
        'parents': [folder_id],
        'mimeType': mime,
        'appProperties': {'cba_key': key},
    }
    created = drive.api.files().create(
        body=body,
        media_body=MediaIoBaseUpload(io.BytesIO(content), mimetype=mime),
        fields='id',
        supportsAllDrives=True,
    ).execute(num_retries=0)
    return drive.meta(created['id']), True


def attachment_spec(value):
    try:
        key, name, path = value.split(':', 2)
    except ValueError as exc:
        raise argparse.ArgumentTypeError('attachment must be KEY:NAME:PATH') from exc
    path = Path(path)
    if not key or not name or not path.is_file():
        raise argparse.ArgumentTypeError(f'invalid attachment: {value}')
    return key, name, path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--candidate', type=Path, required=True)
    parser.add_argument('--acceptance', type=Path, required=True)
    parser.add_argument('--sheet-folder-id', required=True)
    parser.add_argument('--sheet-name', required=True)
    parser.add_argument('--sheet-key', required=True)
    parser.add_argument('--attachment-folder-id')
    parser.add_argument('--attachment', action='append', type=attachment_spec, default=[])
    parser.add_argument('--report', type=Path, required=True)
    parser.add_argument('--root', type=Path, default=Path.cwd(),
                        help='Deployment root that holds .credentials')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()

    local, checks = local_checks(args.candidate, args.acceptance)
    if args.attachment and not args.attachment_folder_id:
        parser.error('--attachment-folder-id is required with --attachment')
    if not args.apply:
        print(json.dumps({
            'status': 'LOCAL_VALIDATED_WITHOUT_WRITE',
            'candidate': str(args.candidate),
            'checks': checks,
        }, ensure_ascii=False, indent=2))
        return

    root = args.root.resolve()
    drive = Drive(root)
    sheet, sheet_created = ensure_sheet(
        drive, args.sheet_folder_id, args.sheet_key, args.sheet_name, args.candidate)
    remote, exported = sheet_values(drive, sheet['id'])
    compare_values(local, remote)

    attachments = []
    for key, name, path in args.attachment:
        mime = 'application/json' if path.suffix.lower() == '.json' else 'text/markdown'
        item, created = ensure_attachment(
            drive, args.attachment_folder_id, key, name, mime, path)
        attachments.append({
            'key': key,
            'id': item['id'],
            'name': item['name'],
            'url': item['webViewLink'],
            'sha256': digest(path.read_bytes()),
            'created': created,
        })

    report = {
        'status': 'DRIVE_CANDIDATE_SYNCED',
        'production_eligible': False,
        'verified_at': datetime.now(timezone.utc).isoformat(),
        'candidate': {
            'sheet_id': sheet['id'],
            'name': sheet['name'],
            'url': sheet['webViewLink'],
            'version': sheet.get('version'),
            'sheet_names': SHEET_ORDER,
            'rows_total': checks['rows_total'],
            'workbook_sha256': checks['workbook_sha256'],
            'drive_export_sha256': digest(exported),
            'semantic_sha256': semantic_sha(remote),
            'created': sheet_created,
        },
        'attachments': attachments,
        'checks': [
            'local workbook matches acceptance report',
            'acceptance report is production_eligible=false',
            'six-tab order and row counts match',
            'Drive six-tab values match every local cell',
            'synced objects carry immutable cba_key app properties',
        ],
    }
    atomic(args.report, (json.dumps(report, ensure_ascii=False, indent=2) + '\n').encode())
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
