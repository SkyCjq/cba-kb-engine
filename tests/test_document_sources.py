import io
import json
from pathlib import Path

import pytest

from cba_kb.document_sources import (
    ReviewRequired,
    parse_docx,
    parse_drive_document,
    parse_ima,
    parse_markdown,
    parse_pdf,
    parse_txt,
)


PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDAT\x08\xd7c\xf8"
    b"\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99=\x1d\x00\x00\x00\x00"
    b"IEND\xaeB`\x82"
)


def text_pdf(value):
    stream = f"BT /F1 12 Tf 72 720 Td ({value}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def test_txt_markdown_and_ima_are_local_only(tmp_path):
    txt = parse_txt(b"\xef\xbb\xbfhello\r\nworld")
    assert txt["text"] == "hello\r\nworld"

    clip = tmp_path / "clip.md"
    image = "https://mmbiz.qpic.cn/example.png"
    raw = (
        "---\n"
        "title: CBA update\n"
        "source_url: https://mp.weixin.qq.com/s/example\n"
        "tags: [外援]\n"
        "attachments: [https://example.com/remote-note.txt]\n"
        "---\n"
        f"body\n\n![image]({image})\n"
    ).encode()
    parsed = parse_markdown(raw, source_locator=clip.name, base_dir=tmp_path)
    assert parsed["title"] == "CBA update"
    assert parsed["canonical_url"] == "https://mp.weixin.qq.com/s/example"
    assert parsed["tags"] == ["外援"]
    assert parsed["remote_references"] == [image]
    assert any(item["kind"] == "remote" and item["url"] == image
               for item in parsed["attachments"])

    ima = parse_ima(raw, source_locator="export.md", base_dir=tmp_path)
    assert ima["capture_channel"] == "ima_file_export"
    assert ima["text"] == parsed["text"]


def test_txt_requires_utf8():
    with pytest.raises(ReviewRequired, match="TXT_ENCODING_UNSUPPORTED"):
        parse_txt(b"\xff\xfe\x00")


def test_docx_preserves_paragraph_table_and_media_order():
    from docx import Document

    document = Document()
    document.add_paragraph("first")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.add_paragraph("last")
    document.add_picture(io.BytesIO(PNG))
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_docx(buffer.getvalue(), source_locator="fixture.docx")
    assert parsed["text"].splitlines()[:3] == ["first", "left\tright", "last"]
    assert parsed["relationship_provenance"]
    assert all(item["relationship_id"].startswith("rId")
               for item in parsed["relationship_provenance"])


def test_pdf_text_and_ocr_required():
    parsed = parse_pdf(text_pdf("CBA PDF"), source_locator="fixture.pdf")
    assert "CBA PDF" in parsed["text"]
    with pytest.raises(ReviewRequired, match="OCR_REQUIRED") as error:
        from pypdf import PdfWriter

        output = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.write(output)
        parse_pdf(output.getvalue(), source_locator="scan.pdf")
    assert error.value.metadata["text_layer"] is False


class FakeDocs:
    def __init__(self, document):
        self.document_value = document
        self.calls = 0

    def document(self, file_id):
        self.calls += 1
        return self.document_value


class FakeDrive:
    def __init__(self, document):
        self.docs = FakeDocs(document)
        self.meta_calls = 0
        self.reset_calls = 0
        self.writes = 0

    def meta(self, file_id):
        self.meta_calls += 1
        return {
            "id": file_id,
            "name": "Fixture Doc",
            "mimeType": "application/vnd.google-apps.document",
            "webViewLink": f"https://docs.google.com/document/d/{file_id}/edit",
        }

    def _reset_read_connections(self):
        self.reset_calls += 1


def paragraph(value):
    return {
        "paragraph": {
            "elements": [{"textRun": {"content": value}}],
        }
    }


def google_document():
    child = {
        "tabProperties": {"tabId": "child", "title": "Child", "index": 1},
        "documentTab": {"body": {"content": [paragraph("child body\n")]}},
        "childTabs": [],
    }
    first = {
        "tabProperties": {"tabId": "first", "title": "First", "index": 2},
        "documentTab": {
            "body": {
                "content": [
                    paragraph("first body\n"),
                    {
                        "table": {
                            "tableRows": [{
                                "tableCells": [
                                    {"body": {"content": [paragraph("A")]}},
                                    {"body": {"content": [paragraph("B")]}},
                                ]
                            }]
                        }
                    },
                ]
            }
        },
        "childTabs": [],
    }
    parent = {
        "tabProperties": {"tabId": "parent", "title": "Parent", "index": 1},
        "documentTab": {"body": {"content": [paragraph("parent body\n")]}},
        "childTabs": [child],
    }
    return {
        "documentId": "public-fixture",
        "revisionId": "revision-1",
        "tabs": [first, parent],
    }


def test_google_docs_preserves_all_tabs_and_topology():
    drive = FakeDrive(google_document())
    parsed = parse_drive_document(drive, "public-fixture")
    assert drive.meta_calls == 1
    assert drive.docs.calls == 1
    assert drive.writes == 0
    assert "parent body" in parsed["text"]
    assert "child body" in parsed["text"]
    assert "A\tB" in parsed["text"]
    assert parsed["tab_topology"] == [
        {
            "tab_id": "parent",
            "title": "Parent",
            "parent_tab_id": None,
            "child_tab_ids": ["child"],
        },
        {
            "tab_id": "child",
            "title": "Child",
            "parent_tab_id": "parent",
            "child_tab_ids": [],
        },
        {
            "tab_id": "first",
            "title": "First",
            "parent_tab_id": None,
            "child_tab_ids": [],
        },
    ]
    raw = json.loads(parsed["raw"])
    assert raw["revisionId"] == "revision-1"


@pytest.mark.parametrize("document", [
    {},
    {"tabs": []},
    {"tabs": [{"tabProperties": {"tabId": "x"}}]},
    {"tabs": [{"tabProperties": {"tabId": "x"}, "documentTab": {}}]},
])
def test_google_docs_partial_structures_fail_closed(document):
    with pytest.raises(ReviewRequired):
        parse_drive_document(FakeDrive(document), "public-fixture")
