import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cba_kb.identity_coverage import prepare_review_packet
from cba_kb.identity_web_evidence import (
    IdentityWebEvidenceError,
    batch_predicate,
    build_batch_plan,
    build_conflicts,
    build_collector_search_manifest,
    build_discovery_record,
    build_evidence_item,
    build_evidence_manifest,
    build_groups,
    classify_record_evidence,
    classify_candidate_evidence,
    collect_evidence_from_responses,
    deterministic_group_id,
    expand_batch_event,
    expand_group_event,
    export_review_workbook,
    external_person_id_claim,
    fetch_official_resource,
    extract_html_claims,
    import_review_workbook,
    make_claim,
    normalize_claim,
    search_exhaustion_status,
    validate_fetch_request,
    validate_redirect,
    validate_source_url,
)


def proposal(record_key, proposal_type, relation, **overrides):
    value = {
        "record_key": record_key,
        "proposal_type": proposal_type,
        "candidate_group_id": None,
        "candidate_player_uid": None,
        "candidate_name_display_only": "Synthetic",
        "proposed_relation": relation,
        "machine_suggestion": relation,
        "machine_reason": "synthetic",
        "evidence_refs": ["synthetic-evidence"],
    }
    value.update(overrides)
    return value


def evidence_item(
    url,
    content,
    *,
    tier="A1",
    claims=None,
    record_keys=None,
    content_type="text/html",
):
    return build_evidence_item(
        source_tier=tier,
        source_kind="registration_notice",
        source_url=url,
        fetched_at="2026-09-18T00:00:00Z",
        http_status=200,
        content_type=content_type,
        content=content,
        raw_snapshot_path=f"raw/{hash(content)}.bin",
        claims=claims or [],
        record_keys=record_keys or [],
    )


def test_source_policy_and_redirect_guard():
    assert validate_source_url(
        "https://www.cba.net.cn/notice",
        source_tier="A0",
    )["domain"] == "www.cba.net.cn"
    with pytest.raises(IdentityWebEvidenceError, match="DOMAIN_NOT_ALLOWED"):
        validate_source_url(
            "https://evil.example/notice",
            source_tier="A0",
        )
    with pytest.raises(IdentityWebEvidenceError, match="B0_DOMAIN_NOT_APPROVED"):
        validate_source_url(
            "https://club.example/notice",
            source_tier="B0",
            private_b0_domains=["official-club.example"],
        )
    with pytest.raises(IdentityWebEvidenceError, match="DOMAIN_NOT_ALLOWED"):
        validate_redirect(
            "https://www.cbaleague.com/a",
            "https://evil.example/b",
            source_tier="A1",
        )


def test_fetch_request_rejects_authenticated_paths_and_headers():
    with pytest.raises(IdentityWebEvidenceError, match="AUTHENTICATED_HEADERS"):
        validate_fetch_request(
            "https://www.cbaleague.com/player",
            headers={"Coo" + "kie": "1"},
        )
    with pytest.raises(IdentityWebEvidenceError, match="AUTHENTICATED_RESOURCE"):
        validate_source_url(
            "https://www.cbaleague.com/login",
            source_tier="A1",
        )
    assert validate_fetch_request(
        "https://www.cbaleague.com/player",
    )["method"] == "GET"


def test_external_person_id_requires_semantics_and_namespace():
    claim = external_person_id_claim(
        namespace="CBA_OFFICIAL_PLAYER_ID",
        identifier="123",
        source_semantics="PLAYER_ENTITY",
        source_locator="json.player.id",
    )
    assert claim["normalized_value"] == "CBA_OFFICIAL_PLAYER_ID:123"
    with pytest.raises(IdentityWebEvidenceError, match="SEMANTICS_REQUIRED"):
        external_person_id_claim(
            namespace="CBA_OFFICIAL_PLAYER_ID",
            identifier="123",
            source_semantics="ARTICLE_ID",
            source_locator="article.id",
        )


def test_evidence_ids_are_deterministic_and_change_with_bytes():
    first = evidence_item(
        "https://www.cbaleague.com/player/1",
        b"<html>one</html>",
    )
    second = evidence_item(
        "https://www.cbaleague.com/player/1",
        b"<html>one</html>",
    )
    changed = evidence_item(
        "https://www.cbaleague.com/player/1",
        b"<html>two</html>",
    )
    assert first["evidence_id"] == second["evidence_id"]
    assert changed["evidence_id"] != first["evidence_id"]


def test_html_claim_extraction_and_normalization():
    claims = extract_html_claims(
        b'<span data-claim-type="OFFICIAL_PLAYER_NAME" '
        b'data-claim-value="  Synthetic  Player  " '
        b'data-claim-locator="player-name"></span>'
    )
    assert claims[0]["normalized_value"] == "Synthetic Player"
    assert normalize_claim("OFFICIAL_BIRTH_DATE", "2000/01/02") == "2000-01-02"


def test_evidence_classification_w1_w2_w3_w4_wx():
    w1 = build_evidence_manifest([
        evidence_item(
            "https://www.cbaleague.com/player/1",
            b"w1",
            claims=[external_person_id_claim(
                namespace="CBA_OFFICIAL_PLAYER_ID",
                identifier="1",
                source_semantics="PLAYER_ENTITY",
                source_locator="player.id",
            )],
            record_keys=["r1"],
        )
    ])
    assert classify_record_evidence("r1", w1)[0] == (
        "W1_OFFICIAL_PERSON_ID_EXACT"
    )

    w2 = build_evidence_manifest([
        evidence_item(
            "https://www.cbaleague.com/player/1",
            b"w2a",
            claims=[
                make_claim(
                    claim_type="OFFICIAL_PLAYER_NAME",
                    raw_value="Synthetic Player",
                    source_locator="name",
                ),
                make_claim(
                    claim_type="OFFICIAL_BIRTH_DATE",
                    raw_value="2000-01-02",
                    source_locator="dob",
                ),
            ],
            record_keys=["r2"],
        ),
        evidence_item(
            "https://www.cbaleague.com/player/2",
            b"w2b",
            claims=[
                make_claim(
                    claim_type="OFFICIAL_PLAYER_NAME",
                    raw_value="Synthetic Player",
                    source_locator="name",
                ),
                make_claim(
                    claim_type="OFFICIAL_TEAM",
                    raw_value="Synthetic Club",
                    source_locator="team",
                ),
            ],
            record_keys=["r2"],
        ),
    ])
    assert classify_record_evidence("r2", w2)[0] == (
        "W2_OFFICIAL_BIO_MULTI_SOURCE"
    )

    w3 = build_evidence_manifest([
        evidence_item(
            "https://www.cbaleague.com/player/3",
            b"w3",
            claims=[
                make_claim(
                    claim_type="OFFICIAL_PLAYER_NAME",
                    raw_value="Synthetic Player",
                    source_locator="name",
                ),
                make_claim(
                    claim_type="OFFICIAL_TEAM",
                    raw_value="Synthetic Club",
                    source_locator="team",
                ),
            ],
            record_keys=["r3"],
        )
    ])
    assert classify_record_evidence("r3", w3)[0] == (
        "W3_OFFICIAL_CONTINUITY"
    )
    w4 = build_evidence_manifest([
        evidence_item(
            "https://www.cbaleague.com/discovery",
            b"w4",
            record_keys=["r4"],
        )
    ])
    assert classify_record_evidence("r4", w4)[0] == (
        "W4_DISCOVERY_SUPPORT"
    )

    wx = build_evidence_manifest([
        evidence_item(
            "https://www.cbaleague.com/player/4",
            b"wx-a",
            claims=[external_person_id_claim(
                namespace="CBA_OFFICIAL_PLAYER_ID",
                identifier="4",
                source_semantics="PLAYER_ENTITY",
                source_locator="id",
            )],
            record_keys=["r5"],
        ),
        evidence_item(
            "https://www.cbaleague.com/player/5",
            b"wx-b",
            claims=[external_person_id_claim(
                namespace="CBA_OFFICIAL_PLAYER_ID",
                identifier="5",
                source_semantics="PLAYER_ENTITY",
                source_locator="id",
            )],
            record_keys=["r5"],
        ),
    ])
    assert classify_record_evidence("r5", wx)[0] == "WX_CONFLICT"
    assert build_conflicts(wx)[0]["record_key"] == "r5"


def test_search_exhaustion_requires_complete_clean_collectors():
    complete = build_collector_search_manifest(
        record_key="r1",
        applicable_required_collectors=["A1"],
        completed_collectors=["A1"],
        failed_collectors=[],
        attempted_official_urls=["https://www.cbaleague.com/player/1"],
        collector_version="v1",
        text_warning=False,
        identity_conflict=False,
        surviving_candidate_count=0,
    )
    assert search_exhaustion_status(
        complete,
    ) == "OFFICIAL_SEARCH_EXHAUSTED_NO_SAFE_MATCH"
    no_collectors = build_collector_search_manifest(
        record_key="r1",
        applicable_required_collectors=[],
        completed_collectors=[],
        failed_collectors=[],
        attempted_official_urls=[],
        collector_version="v1",
        text_warning=False,
        identity_conflict=False,
        surviving_candidate_count=0,
    )
    assert search_exhaustion_status(
        no_collectors,
    ) != "OFFICIAL_SEARCH_EXHAUSTED_NO_SAFE_MATCH"
    failed = build_collector_search_manifest(
        record_key="r1",
        applicable_required_collectors=["A1"],
        completed_collectors=[],
        failed_collectors=["A1"],
        attempted_official_urls=["https://www.cbaleague.com/player/1"],
        collector_version="v1",
        text_warning=False,
        identity_conflict=False,
        surviving_candidate_count=0,
    )
    assert search_exhaustion_status(
        failed,
    ) != "OFFICIAL_SEARCH_EXHAUSTED_NO_SAFE_MATCH"


def test_group_ids_deterministic_and_group_event_exact_membership():
    members = ["review_b", "review_a"]
    assert deterministic_group_id(
        members,
        "W1_OFFICIAL_PERSON_ID_EXACT",
    ) == deterministic_group_id(
        reversed(members),
        "W1_OFFICIAL_PERSON_ID_EXACT",
    )
    packet = prepare_review_packet([
        proposal(
            "r1",
            "NEW_IDENTITY_CANDIDATE",
            "PROPOSED_SAME",
            candidate_group_id="group-1",
        ),
        proposal(
            "r2",
            "NEW_IDENTITY_CANDIDATE",
            "PROPOSED_SAME",
            candidate_group_id="group-1",
        ),
    ])
    groups = [{
        "candidate_group_id": "group-1",
        "member_review_ids": sorted(
            item["review_id"] for item in packet["reviews"]
        ),
        "member_count": 2,
        "evidence_class": "W1_OFFICIAL_PERSON_ID_EXACT",
        "conflict_count": 0,
    }]
    manifest = build_evidence_manifest([])
    event = {
        "event_id": "event-1",
        "candidate_group_id": "group-1",
        "review_packet_sha256": packet["review_packet_sha256"],
        "web_evidence_manifest_sha256": manifest[
            "web_evidence_manifest_sha256"
        ],
        "member_review_ids": groups[0]["member_review_ids"],
        "human_decision": "APPROVE",
        "human_note": "approved group",
        "approved_player_uid": "pid_0000000000000099",
        "approved_canonical_name": "Synthetic Player",
        "allocation_attestation": "allocated independently",
        "reviewed_at": "2026-09-18T00:00:00Z",
    }
    expanded = expand_group_event(packet, manifest, groups, event)
    assert len(expanded) == 2
    assert all(item["approved_player_uid"] == "pid_0000000000000099" for item in expanded)
    with pytest.raises(IdentityWebEvidenceError, match="MEMBERSHIP_MISMATCH"):
        expand_group_event(
            packet,
            manifest,
            groups,
            {**event, "member_review_ids": groups[0]["member_review_ids"][:1]},
        )


def test_batch_predicate_and_event_expansion():
    assert batch_predicate({
        "proposal_type": "EXISTING_IDENTITY_CANDIDATE",
        "evidence_class": "W1_OFFICIAL_PERSON_ID_EXACT",
        "conflict_count": 0,
        "text_corruption_warning": False,
    }) == "BATCH_EXISTING_W1"
    assert batch_predicate({
        "proposal_type": "EXISTING_IDENTITY_CANDIDATE",
        "evidence_class": "W2_OFFICIAL_BIO_MULTI_SOURCE",
        "conflict_count": 1,
        "text_corruption_warning": False,
    }) is None
    packet = prepare_review_packet([
        proposal(
            "r1",
            "EXISTING_IDENTITY_CANDIDATE",
            "KEEP_UNDECIDED",
            candidate_player_uid="pid_0000000000000001",
        )
    ])
    review_id = packet["reviews"][0]["review_id"]
    batches = [{
        "batch_id": "batch-1",
        "batch_predicate": "BATCH_EXISTING_W1",
        "member_review_ids": [review_id],
        "member_count": 1,
    }]
    manifest = build_evidence_manifest([])
    event = {
        "event_id": "batch-event-1",
        "batch_id": "batch-1",
        "review_packet_sha256": packet["review_packet_sha256"],
        "web_evidence_manifest_sha256": manifest[
            "web_evidence_manifest_sha256"
        ],
        "member_review_ids": [review_id],
        "batch_predicate": "BATCH_EXISTING_W1",
        "human_decision": "APPROVE",
        "human_note": "approved batch",
        "reviewed_at": "2026-09-18T00:00:00Z",
    }
    expanded = expand_batch_event(packet, manifest, batches, event)
    assert "authority_event_id=batch-event-1" in expanded[0]["human_note"]
    assert expanded[0]["approved_player_uid"] is None


def test_batch_plan_partitions_predicates_before_chunking():
    packet = prepare_review_packet([
        proposal(
            "r1",
            "EXISTING_IDENTITY_CANDIDATE",
            "KEEP_UNDECIDED",
            candidate_player_uid="pid_0000000000000001",
        ),
        proposal(
            "r2",
            "NO_SAFE_CANDIDATE",
            "NO_SAFE_CANDIDATE",
            machine_suggestion="NO_SAFE_CANDIDATE",
        ),
    ])
    item = evidence_item(
        "https://www.cbaleague.com/player/1",
        b"batch-bound",
        claims=[external_person_id_claim(
            namespace="CBA_OFFICIAL_PLAYER_ID",
            identifier="1",
            source_semantics="PLAYER_ENTITY",
            source_locator="id",
        )],
        record_keys=["r1"],
    )
    item["bindings"] = [{
        "record_key": "r1",
        "target_type": "EXISTING_UID",
        "target_id": "pid_0000000000000001",
        "matched_claim_fields": ["OFFICIAL_SOURCE_DECLARED_PERSON_ID"],
        "binding_rationale": "explicit player binding",
    }]
    manifest = build_evidence_manifest([item])
    exhausted = build_collector_search_manifest(
        record_key="r2",
        applicable_required_collectors=["A1"],
        completed_collectors=["A1"],
        failed_collectors=[],
        attempted_official_urls=["https://www.cbaleague.com/search"],
        collector_version="v1",
        text_warning=False,
        identity_conflict=False,
        surviving_candidate_count=0,
    )
    plan = build_batch_plan(
        packet,
        manifest,
        batch_size=1,
        search_statuses={"r2": exhausted},
    )
    assert {item["batch_predicate"] for item in plan["batches"]} == {
        "BATCH_EXISTING_W1",
        "BATCH_NO_SAFE_SEARCH_EXHAUSTED",
    }


def test_collect_html_and_pdf_fixtures(tmp_path):
    discovery = build_discovery_record(
        query="Synthetic Player official",
        discovered_url="https://www.cbaleague.com/player/1",
        discovery_provider="search_engine",
        discovered_at="2026-09-18T00:00:00Z",
        title_hint="Synthetic",
    )
    manifest = collect_evidence_from_responses(
        [{
            "discovery": discovery,
            "status": 200,
            "content_type": "text/html",
            "body": (
                b'<span data-claim-type="OFFICIAL_PLAYER_NAME" '
                b'data-claim-value="Synthetic Player"></span>'
            ),
            "record_keys": ["r1"],
        }],
        output_root=tmp_path,
        extracted_at="2026-09-18T00:00:00Z",
        source_tier="A1",
        source_kind="registration_notice",
    )
    assert len(manifest["items"]) == 1
    assert manifest["items"][0]["claims"][0]["normalized_value"] == (
        "Synthetic Player"
    )


def test_fetch_boundary_retry_redirect_timeout_and_size():
    calls = []

    def transport(url, *, method, timeout, headers):
        calls.append((url, method, timeout, headers))
        return {
            "status": 200,
            "content_type": "text/html",
            "body": b"<html>ok</html>",
            "redirects": ["https://www.cbaleague.com/player/1"],
        }

    result = fetch_official_resource(
        "https://www.cbaleague.com/start",
        source_tier="A1",
        transport=transport,
        max_retries=0,
        per_domain_interval=0,
        timeout_seconds=3,
    )
    assert result["content"] == b"<html>ok</html>"
    assert calls[0][1] == "GET"

    def oversized(url, *, method, timeout, headers):
        return {
            "status": 200,
            "content_type": "text/html",
            "body": b"x" * 10,
            "redirects": [],
        }

    with pytest.raises(IdentityWebEvidenceError, match="TOO_LARGE"):
        fetch_official_resource(
            "https://www.cbaleague.com/start",
            source_tier="A1",
            transport=oversized,
            max_retries=0,
            per_domain_interval=0,
            max_bytes=5,
        )


def test_candidate_bound_evidence_blocks_wrong_attachment():
    item = evidence_item(
        "https://www.cbaleague.com/player/1",
        b"bound",
        claims=[external_person_id_claim(
            namespace="CBA_OFFICIAL_PLAYER_ID",
            identifier="1",
            source_semantics="PLAYER_ENTITY",
            source_locator="id",
        )],
        record_keys=["r1"],
    )
    item["bindings"] = [{
        "record_key": "r1",
        "target_type": "EXISTING_UID",
        "target_id": "pid_0000000000000001",
        "matched_claim_fields": ["OFFICIAL_SOURCE_DECLARED_PERSON_ID"],
        "binding_rationale": "explicit official player entity binding",
    }]
    manifest = build_evidence_manifest([item])
    assert classify_candidate_evidence(
        "r1",
        target_type="EXISTING_UID",
        target_id="pid_0000000000000001",
        manifest=manifest,
    )[0] == "W1_OFFICIAL_PERSON_ID_EXACT"
    assert classify_candidate_evidence(
        "r1",
        target_type="EXISTING_UID",
        target_id="pid_0000000000000002",
        manifest=manifest,
    )[0] == "W4_DISCOVERY_SUPPORT"


def test_workbook_export_import_unicode_and_machine_tamper(tmp_path):
    packet = prepare_review_packet([
        proposal(
            "r1",
            "EXISTING_IDENTITY_CANDIDATE",
            "KEEP_UNDECIDED",
            candidate_player_uid="pid_0000000000000001",
            candidate_name_display_only="中文球员",
        )
    ])
    manifest = build_evidence_manifest([])
    workbook_path = tmp_path / "review_packet_workbook.xlsx"
    export_review_workbook(
        workbook_path,
        packet=packet,
        evidence_manifest=manifest,
        groups=[],
        batches=[],
    )
    workbook = load_workbook(workbook_path)
    assert workbook["Review"]["F2"].value == "中文球员"
    review_id = packet["reviews"][0]["review_id"]
    review_sheet = workbook["Review"]
    review_sheet["K2"] = "UNDECIDED"
    review_sheet["L2"] = "reviewed by synthetic human"
    review_sheet["P2"] = "2026-09-18T00:00:00Z"
    workbook.save(workbook_path)
    csv_bytes = import_review_workbook(
        workbook_path,
        packet=packet,
        evidence_manifest=manifest,
        groups=[],
        batches=[],
    )
    assert review_id.encode() in csv_bytes

    tampered = load_workbook(workbook_path)
    tampered["Review"]["B2"] = "tampered"
    tampered.save(workbook_path)
    with pytest.raises(IdentityWebEvidenceError, match="MACHINE_FIELD_TAMPER"):
        import_review_workbook(
            workbook_path,
            packet=packet,
            evidence_manifest=manifest,
            groups=[],
            batches=[],
        )


def test_workbook_rejects_group_batch_and_evidence_machine_edits(tmp_path):
    packet = prepare_review_packet([
        proposal(
            "r1",
            "NEW_IDENTITY_CANDIDATE",
            "PROPOSED_SAME",
            candidate_group_id="group-1",
        )
    ])
    evidence = evidence_item(
        "https://www.cbaleague.com/player/1",
        b"workbook-evidence",
        record_keys=["r1"],
        claims=[external_person_id_claim(
            namespace="CBA_OFFICIAL_PLAYER_ID",
            identifier="1",
            source_semantics="PLAYER_ENTITY",
            source_locator="id",
        )],
    )
    manifest = build_evidence_manifest([evidence])
    review_id = packet["reviews"][0]["review_id"]
    groups = [{
        "candidate_group_id": "group-1",
        "member_review_ids": [review_id],
        "member_count": 1,
        "evidence_class": "W1_OFFICIAL_PERSON_ID_EXACT",
        "conflict_count": 0,
    }]
    batches = [{
        "batch_id": "batch-1",
        "batch_predicate": "BATCH_EXISTING_W1",
        "member_review_ids": [review_id],
        "member_count": 1,
    }]
    workbook_path = tmp_path / "review_packet_workbook.xlsx"
    export_review_workbook(
        workbook_path,
        packet=packet,
        evidence_manifest=manifest,
        groups=groups,
        batches=batches,
    )
    workbook = load_workbook(workbook_path)
    workbook["Groups"]["C2"] = 2
    workbook.save(workbook_path)
    with pytest.raises(IdentityWebEvidenceError, match="Groups.member_count"):
        import_review_workbook(
            workbook_path,
            packet=packet,
            evidence_manifest=manifest,
            groups=groups,
            batches=batches,
        )

    export_review_workbook(
        workbook_path,
        packet=packet,
        evidence_manifest=manifest,
        groups=groups,
        batches=batches,
    )
    workbook = load_workbook(workbook_path)
    workbook["Batches"]["B2"] = "BATCH_EXISTING_W2"
    workbook.save(workbook_path)
    with pytest.raises(IdentityWebEvidenceError, match="Batches.batch_predicate"):
        import_review_workbook(
            workbook_path,
            packet=packet,
            evidence_manifest=manifest,
            groups=groups,
            batches=batches,
        )

    export_review_workbook(
        workbook_path,
        packet=packet,
        evidence_manifest=manifest,
        groups=groups,
        batches=batches,
    )
    workbook = load_workbook(workbook_path)
    workbook["Evidence_Index"]["C2"] = "https://evil.example/"
    workbook.save(workbook_path)
    with pytest.raises(IdentityWebEvidenceError, match="Evidence_Index.source_url"):
        import_review_workbook(
            workbook_path,
            packet=packet,
            evidence_manifest=manifest,
            groups=groups,
            batches=batches,
        )
